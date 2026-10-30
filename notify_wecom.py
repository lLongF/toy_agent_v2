#!/usr/bin/env python3
"""把测试结果 Excel 发到企业微信群机器人。"""
from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path
from typing import Optional, Tuple
from urllib.parse import parse_qs, urlparse

import requests

DEFAULT_WEBHOOK = os.environ.get("WECOM_WEBHOOK_URL", "").strip()


def webhook_key(url: str) -> str:
    qs = parse_qs(urlparse(url).query)
    keys = qs.get("key") or []
    if not keys or not keys[0]:
        raise ValueError("Webhook 地址里没有 key 参数")
    return keys[0]


def xlsx_detailed_stats(path: Path):
    """返回 (总行数, 通过, 业务不通过, 评分失败, 未执行)。"""
    try:
        import openpyxl
    except ImportError:
        return 0, 0, 0, 0, 0
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Test case"] if "Test case" in wb.sheetnames else wb.active
        headers = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col = None
        for i, h in enumerate(headers):
            if h == "测试结果是否通过":
                col = i
                break
        if col is None:
            return 0, 0, 0, 0, 0
        total = passed = failed = judge_failed = unexecuted = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue
            total += 1
            val = str(row[col]).strip() if col < len(row) and row[col] is not None else ""
            if val == "通过":
                passed += 1
            elif val == "评分失败":
                judge_failed += 1
            elif val in ("未执行", ""):
                unexecuted += 1
            else:
                failed += 1
        return total, passed, failed, judge_failed, unexecuted
    finally:
        wb.close()


def xlsx_stats(path: Path) -> Tuple[int, int, int, int]:
    """兼容原四元组接口；未执行不混入业务失败。"""
    return xlsx_detailed_stats(path)[:4]


def send_text(webhook: str, content: str) -> None:
    r = requests.post(
        webhook,
        json={"msgtype": "text", "text": {"content": content[:2048]}},
        timeout=30,
    )
    r.raise_for_status()
    data = r.json()
    if data.get("errcode") not in (0, None):
        raise RuntimeError(f"企微文本发送失败: {data}")


def send_file(webhook: str, path: Path) -> None:
    key = webhook_key(webhook)
    upload = f"https://qyapi.weixin.qq.com/cgi-bin/webhook/upload_media?key={key}&type=file"
    with path.open("rb") as f:
        r = requests.post(upload, files={"media": (path.name, f)}, timeout=60)
    r.raise_for_status()
    data = r.json()
    media_id = data.get("media_id")
    if not media_id:
        raise RuntimeError(f"企微上传文件失败: {data}")
    r2 = requests.post(
        webhook,
        json={"msgtype": "file", "file": {"media_id": media_id}},
        timeout=30,
    )
    r2.raise_for_status()
    data2 = r2.json()
    if data2.get("errcode") not in (0, None):
        raise RuntimeError(f"企微文件发送失败: {data2}")


def notify(xlsx: Optional[Path], webhook: str, extra: str = "") -> None:
    lines = ["玩具助手每日文本测试"]
    if extra:
        lines.append(extra.strip())
    if xlsx and xlsx.is_file():
        total, passed, failed, judge_failed, unexecuted = xlsx_detailed_stats(xlsx)
        if total:
            lines.append(f"用例 {total} 条，通过 {passed}，不通过 {failed}，评分失败 {judge_failed}，未执行 {unexecuted}")
        lines.append(f"附件：{xlsx.name}")
    else:
        lines.append("没有生成结果 Excel。")
    send_text(webhook, "\n".join(lines))
    if xlsx and xlsx.is_file():
        send_file(webhook, xlsx)


def main() -> int:
    parser = argparse.ArgumentParser(description="发送测试结果到企微群")
    parser.add_argument("xlsx", nargs="?", help="结果 xlsx 路径")
    parser.add_argument("--webhook", default=DEFAULT_WEBHOOK, help="企微机器人 Webhook，默认读环境变量 WECOM_WEBHOOK_URL")
    parser.add_argument("--message", default="", help="附加说明，例如失败原因")
    args = parser.parse_args()
    if not args.webhook:
        print("未配置 WECOM_WEBHOOK_URL，跳过企微通知。", file=sys.stderr)
        return 2
    xlsx = Path(args.xlsx) if args.xlsx else None
    notify(xlsx, args.webhook, args.message)
    print("已发送到企微群。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
