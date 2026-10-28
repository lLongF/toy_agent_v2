"""Small JSON dataset registry and immutable per-run workbook snapshots."""
from __future__ import annotations

import json
import os
import shutil
from dataclasses import asdict, dataclass, replace
from pathlib import Path
from uuid import uuid4

import openpyxl

from config import _dialogue_worksheet, load_dialogue_groups_from_xlsx
from toy_profiles import PROJECT_ROOT, load_toy_profile, safe_path, validate_id
from toy_management import atomic_json, file_hash

HEADERS = ("类型", "用例编号", "消息列表", "预期结果", "是否工具调用", "工具名称", "工具参数文字描述")


@dataclass
class Dataset:
    id: str
    name: str
    scope: str
    owner_toy_id: str
    path: Path
    enabled: bool = True
    description: str = ""


def _root(toy=None, root=None):
    return Path(root or (toy.root.parent.parent if toy is not None else PROJECT_ROOT)).resolve()


def registry_path(root=None):
    return safe_path(_root(root=root) / "settings" / "datasets.json", _root(root=root))


def registry_hash(root=None):
    return file_hash(registry_path(root))


def _validate(dataset, root):
    validate_id(dataset.id)
    if not isinstance(dataset.name, str) or not dataset.name.strip():
        raise ValueError("测试集名称不能为空")
    if dataset.scope not in ("public", "toy") or not isinstance(dataset.enabled, bool):
        raise ValueError("测试集范围或状态无效")
    if dataset.scope == "toy":
        validate_id(dataset.owner_toy_id)
    elif dataset.owner_toy_id:
        raise ValueError("公共测试集不能指定所属玩具")
    path = safe_path(root / dataset.path, root)
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        raise ValueError("测试集必须是项目内已有 xlsx")
    return replace(dataset, path=path)


def load_registry(root=None):
    root = _root(root=root)
    path = registry_path(root)
    if not path.exists():
        return []
    raw = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict) or raw.get("version") != 1 or not isinstance(raw.get("datasets"), list):
        raise ValueError("测试集注册表格式无效")
    datasets = [_validate(Dataset(**item), root) for item in raw["datasets"]]
    if len({d.id.casefold() for d in datasets}) != len(datasets):
        raise ValueError("测试集 id 重复")
    return datasets


def save_registry(datasets, expected_hash, root=None):
    root = _root(root=root)
    checked = [_validate(d, root) for d in datasets]
    if len({d.id.casefold() for d in checked}) != len(checked):
        raise ValueError("测试集 id 重复")
    if any(d.id.startswith("legacy-") for d in checked):
        raise ValueError("legacy- 为兼容测试集保留前缀")
    output = []
    for dataset in checked:
        item = asdict(dataset)
        item["path"] = dataset.path.relative_to(root).as_posix()
        output.append(item)
    return atomic_json(registry_path(root), {"version": 1, "datasets": output}, expected_hash, root)


def list_datasets(toy, root=None):
    """List public/owned entries including disabled ones; no migration writes."""
    root = _root(toy, root)
    entries = [d for d in load_registry(root) if d.scope == "public" or d.owner_toy_id == toy.id]
    legacy_path = safe_path(toy.cases_xlsx, root)
    if legacy_path.is_file() and not any(d.path == legacy_path for d in entries):
        entries.insert(0, Dataset("legacy-" + toy.id, (toy.name or toy.id) + " 默认测试集",
                                  "toy", toy.id, legacy_path))
    return entries


def resolve_datasets(toy, selected_ids, root=None):
    if not toy.enabled:
        raise ValueError("玩具已禁用")
    entries = list_datasets(toy, root)
    ids = list(selected_ids or [])
    if not ids:
        # The historical file remains the default even after a registry is introduced.
        match = next((d for d in entries if d.path == Path(toy.cases_xlsx).resolve()), None)
        if match is None:
            raise ValueError("默认测试集不存在，请显式选择测试集")
        ids = [match.id]
    if len(set(ids)) != len(ids):
        raise ValueError("重复选择测试集")
    available = {d.id: d for d in entries}
    result = []
    for identity in ids:
        if identity not in available:
            raise ValueError("测试集不存在或不属于该玩具")
        dataset = available[identity]
        if not dataset.enabled:
            raise ValueError("测试集已禁用")
        result.append(dataset)
    return result


def create_dataset(name, scope="public", owner_toy_id="", source_path=None, *,
                   dataset_id=None, enabled=True, description="", expected_hash=None, root=None):
    """Copy a project XLSX or create a header-only workbook; register with CAS."""
    root = _root(root=root)
    identity = validate_id(dataset_id or "dataset-" + uuid4().hex[:12])
    if identity.startswith("legacy-"):
        raise ValueError("保留的测试集 id")
    current_hash = registry_hash(root)
    if expected_hash is not None and expected_hash != current_hash:
        raise ValueError("配置已修改，请重新加载后保存")
    entries = load_registry(root)
    if any(d.id.casefold() == identity.casefold() for d in entries):
        raise ValueError("测试集 id 重复")
    target = safe_path(root / "data" / "datasets" / (identity + ".xlsx"), root)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = safe_path(target.parent / ("." + uuid4().hex + ".xlsx"), root)
    dataset = Dataset(identity, name, scope, owner_toy_id, target, enabled, description)
    # Validate metadata before creating a file (path existence is checked after publishing).
    if scope not in ("public", "toy") or not isinstance(enabled, bool) or not str(name).strip():
        raise ValueError("测试集名称、范围或状态无效")
    if scope == "toy":
        validate_id(owner_toy_id)
        load_toy_profile(owner_toy_id, root / "toys")
    elif owner_toy_id:
        raise ValueError("公共测试集不能指定所属玩具")
    published = False
    try:
        if source_path is not None:
            source = safe_path(root / Path(source_path), root)
            if source.suffix.lower() != ".xlsx" or not source.is_file():
                raise ValueError("来源必须是项目内已有 xlsx")
            load_dialogue_groups_from_xlsx(source)
            shutil.copy2(source, temporary)
        else:
            workbook = openpyxl.Workbook()
            workbook.active.title = "Test case"
            workbook.active.append(HEADERS)
            workbook.save(temporary)
            workbook.close()
        # Exclusive creation prevents collisions from replacing any existing data.
        with target.open("xb") as stream, temporary.open("rb") as source:
            shutil.copyfileobj(source, stream)
            stream.flush()
            os.fsync(stream.fileno())
        published = True
        save_registry(entries + [dataset], current_hash, root)
    except Exception:
        if published:
            # This invocation exclusively created this new file; no old asset is removed.
            target.unlink()
        raise
    finally:
        temporary.unlink(missing_ok=True)
    return dataset


def prepare_dataset_snapshot(cfg, tempdir, root=None):
    """Copy/merge selected datasets and expose source rows plus isolated sessions.

    rows: {str(execution_row): {dataset_id,dataset_name,sheet,row,category}}
    sessions: [{dataset_id,category,rows:[execution_row,...]}]. Category names stay
    unchanged: the caller MUST use sessions to avoid merging independent memories.
    """
    root = _root(root=root)
    toy = load_toy_profile(cfg.toy_id, root / "toys")
    datasets = resolve_datasets(toy, getattr(cfg, "selected_dataset_ids", []), root)
    tempdir = Path(tempdir).resolve()
    tempdir.mkdir(parents=True, exist_ok=True)
    target = safe_path(tempdir / ("dataset-" + uuid4().hex + ".xlsx"), tempdir)
    snapshot = {"version": 1, "datasets": [], "rows": {}, "sessions": []}
    merged = openpyxl.Workbook()
    merged.active.title = "Test case"
    merged.active.append(HEADERS)
    try:
        for dataset in datasets:
            before = file_hash(dataset.path)
            from asset_store import case_issues
            issues = case_issues(dataset.path)
            if issues:
                raise ValueError("测试集校验失败：" + dataset.name + "：" + "; ".join(issues))
            groups = load_dialogue_groups_from_xlsx(dataset.path)
            if not groups or not any(group.turns for group in groups):
                raise ValueError("空测试集不能执行：" + dataset.name)
            workbook = openpyxl.load_workbook(dataset.path, read_only=True, data_only=True)
            try:
                sheet_name = _dialogue_worksheet(workbook).title
            finally:
                workbook.close()
            snapshot["datasets"].append({"id": dataset.id, "name": dataset.name,
                                           "path": dataset.path.relative_to(root).as_posix(), "sha256": before})
            for group in groups:
                session = {"dataset_id": dataset.id, "category": group.category, "rows": []}
                for turn in group.turns:
                    if len(datasets) == 1:
                        execution_row = turn.sheet_row
                    else:
                        execution_row = merged.active.max_row + 1
                        merged.active.append((group.category, turn.case_id, " / ".join(turn.user_messages),
                                              turn.expected_outcome, "是" if turn.expect_tool_call else "否",
                                              turn.tool_name, turn.tool_parameters))
                    snapshot["rows"][str(execution_row)] = {"dataset_id": dataset.id,
                        "dataset_name": dataset.name, "sheet": sheet_name,
                        "row": turn.sheet_row, "category": group.category}
                    session["rows"].append(execution_row)
                snapshot["sessions"].append(session)
            if len(datasets) == 1:
                shutil.copy2(dataset.path, target)
                if file_hash(target) != before:
                    raise ValueError("制作快照时测试集已变化，请重试")
            if before != file_hash(dataset.path):
                raise ValueError("制作快照时测试集已变化，请重试")
        if len(datasets) > 1:
            merged.save(target)
        load_dialogue_groups_from_xlsx(target)
    except Exception:
        target.unlink(missing_ok=True)
        raise
    finally:
        merged.close()
    return replace(cfg, cases_xlsx=target, case_snapshot=target, dataset_snapshot=snapshot)
