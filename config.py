#!/usr/bin/env python3
"""公共固定配置与用例模型。各玩具只换用例、工具描述、token，不改这里。"""
from __future__ import annotations

import ast
import asyncio
import random
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl

from toy_profiles import DEFAULT_TOY_ID

DEFAULT_DEVICE_MAC = "18:c0:4d:4f:31:3c"
DEFAULT_FIRMWARE_VERSION = "1.0.0"

DEVICE_MAC_LIST = [
    "18:c0:4d:4f:31:3c",
    "30:ed:a0:a6:3d:9d",
    "30:ed:a0:a6:3d:9e",
    "30:ed:a0:a6:3d:9f",
    "30:ed:a0:a6:3d:a0",
    "30:ed:a0:a6:3d:a1",
    "30:ed:a0:a6:3d:a2",
    "30:ed:a0:a6:3d:a3",
    "30:ed:a0:a6:3d:a4",
    "30:ed:a0:a6:3d:a5",
    "30:ed:a0:a6:3d:a6",
]


def generate_random_mac() -> str:
    """生成本地管理、单播 MAC，格式 aa:bb:cc:dd:ee:ff。"""
    first = (random.randint(0, 255) & 0xFE) | 0x02
    rest = [random.randint(0, 255) for _ in range(5)]
    return ":".join(f"{b:02x}" for b in (first, *rest))


# 确保目录存在
def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)

# 工具列表加载
def load_mcp_tools_from_txt(path: Path) -> List[Dict[str, Any]]:
    """从「工具描述.txt」加载 MCP 工具列表（文件内容为 Python 字面量列表元素拼接，用 ast 解析）。"""
    if not path.is_file():
        return []
    raw = path.read_text(encoding="utf-8").strip()
    if not raw:
        return []
    blob = raw if raw.startswith("[") else f"[{raw}]"
    try:
        data = ast.literal_eval(blob)
    except (SyntaxError, ValueError):
        return []
    if not isinstance(data, list):
        return []
    out: List[Dict[str, Any]] = []
    for item in data:
        if not isinstance(item, dict):
            continue

        # 兼容两种常见工具描述格式：
        # 1) MCP：{"name": ..., "description": ..., "inputSchema": ...}
        # 2) OpenAI function：{"type": "function", "function": {"name": ..., "parameters": ...}}
        if item.get("name"):
            name = item.get("name")
            description = item.get("description", "")
            input_schema = item.get("inputSchema")
        else:
            function_def = item.get("function")
            if not isinstance(function_def, dict) or not function_def.get("name"):
                continue
            name = function_def.get("name")
            description = function_def.get("description", "")
            input_schema = function_def.get("parameters")

        out.append(
            {
                "name": str(name),
                "description": str(description),
                "inputSchema": input_schema or {"type": "object", "properties": {}},
            }
        )
    return out

# 无工具描述时的默认工具描述
def default_fallback_mcp_tools() -> List[Dict[str, Any]]:
    return [
        {
            "name": "ping",
            "description": "ping 工具（当「工具描述.txt」缺失或无效时的回退）",
            "inputSchema": {"type": "object", "properties": {}},
        }
    ]


# ---------------------------------------------------------------------------
# Excel 用例模型：与「示例输入.xlsx」列一致；相同「用例编号」归为同一组对话
# ---------------------------------------------------------------------------


@dataclass
class DialogueTurn:
    """表格中的一行；「消息列表」内可用「/」拆成多条连续用户话。"""
    # 单条对话
    sheet_row: int
    case_id: str
    user_messages: List[str]
    expected_outcome: str
    expect_tool_call: Optional[bool]
    tool_name: Optional[str]
    tool_parameters: Optional[str]


@dataclass
class DialogueGroup:
    """同一测试类型下的多轮对话（多行合并为一组，共用一条 WebSocket 会话）。"""
    group_id: str
    category: str
    turns: List[DialogueTurn]

# 是否存在工具调用标识识别
def _parse_tool_flag(cell: Any) -> Optional[bool]:
    if cell is None:
        return None
    s = str(cell).strip()
    if not s:
        return None
    if s in ("是", "yes", "Yes", "YES", "true", "True", "1"):
        return True
    if s in ("否", "no", "No", "NO", "false", "False", "0"):
        return False
    return None

# xlsx中的表头形式为 预期消息/非预期消息 这里在对其中的内容进行拆解
def _split_user_messages(raw: Any) -> List[str]:
    if raw is None:
        return []
    text = str(raw).strip()
    if not text:
        return []
    parts = re.split(r"\s*/\s*", text)
    return [p.strip() for p in parts if p.strip()]

# 文件结构不定 用于匹配新旧结构，后续确定表结构后可删除
def _col_idx_optional(col: Dict[str, int], *names: str) -> Optional[int]:
    for n in names:
        if n in col:
            return col[n]
    return None

# 获取预期调用参数：工具名/工具参数描述
def dialogue_turn_expected_tool(turn: DialogueTurn) -> Optional[str]:
    """MCP 匹配用：优先「工具名称」，否则回退「工具参数 / 工具参数文字描述」。"""
    if turn.tool_name and str(turn.tool_name).strip():
        return str(turn.tool_name).strip()
    return turn.tool_parameters

# 是否工具调用转化：空单元格和「否」都不强制
def turn_requires_mcp(expect_tool_call: Optional[bool]) -> bool:
    """是否要求必须出现 MCP tools/call：仅当「是否工具调用」明确为「是」时才强制。"""
    return expect_tool_call is True

# 判别文本内容拼接
def combined_text_for_mcp_no_tool_llm_judge(turn: Dict[str, Any]) -> str:
    """不要求 MCP 的 turn：给 LLM 裁判的助手侧文本，避免多字段重复拼接。"""
    c = str(turn.get("combined_server_text", "") or "").strip()
    if c:
        return c
    parts: List[str] = []
    for u in turn.get("utterances") or []:
        if not isinstance(u, dict):
            continue
        seg = str(u.get("server_text_for_judge") or u.get("server_text") or "").strip()
        if seg:
            parts.append(seg)
    return "\n".join(parts)


def _dialogue_worksheet(wb: Any) -> Any:
    """与用例读取一致：优先子表「Test case」，否则退回 active（单表工作簿）。"""
    names = getattr(wb, "sheetnames", None) or []
    if "Test case" in names:
        return wb["Test case"]
    return wb.active


def _ensure_test_result_columns(ws: Any, extra_headers: List[str]) -> int:
    """若首行已有「测试结果是否通过」列则复用其起始列，否则在表尾追加表头。返回该组列的起始列号（1-based）。"""
    max_c = int(ws.max_column or 0)
    for c in range(1, max_c + 1):
        v = ws.cell(row=1, column=c).value
        if v is not None and str(v).strip() == extra_headers[0]:
            # 兼容旧版结果表：即使已经存在前四列，也补齐/刷新新增表头。
            for i, title in enumerate(extra_headers):
                ws.cell(row=1, column=c + i, value=title)
            return c
    base = max_c + 1
    for i, title in enumerate(extra_headers):
        ws.cell(row=1, column=base + i, value=title)
    return base


def load_dialogue_groups_from_xlsx(path: Path) -> List[DialogueGroup]:
    """读取 xlsx，按「类型」分组；空「类型」继承上一行有效类型。同一类型按表格顺序共用一条会话。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = _dialogue_worksheet(wb)
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            return []

        headers = [str(h).strip() if h is not None else "" for h in header_row]
        col = {name: i for i, name in enumerate(headers)}

        def col_idx(*names: str) -> int:
            for n in names:
                if n in col:
                    return col[n]
            raise ValueError(f"xlsx 缺少必要列（需其一）: {names}；实际表头: {headers}")
        # 针对新的xlsx表需要进行改动 mcp调用的判别通过工具成功列来判别
        i_type = col_idx("类型")
        i_id = col_idx("用例编号")
        i_msg = col_idx("消息列表")
        i_exp = col_idx("预期结果/非预期结果", "预期结果")
        i_tool = _col_idx_optional(col, "是否工具调用")
        i_tool_name = _col_idx_optional(col, "工具名称")
        # 多了个工具描述列，应该是单纯对工具进行描述
        i_tool_arg = _col_idx_optional(col, "工具参数", "工具参数文字描述")

        last_category = ""
        buckets: Dict[str, List[DialogueTurn]] = {}
        order: List[str] = []

        for row_idx, row in enumerate(rows_iter, start=2):
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue
            cat_cell = row[i_type] if i_type < len(row) else None
            if cat_cell is not None and str(cat_cell).strip():
                last_category = str(cat_cell).strip()
            category = last_category
            if not category:
                continue
            gid_raw = row[i_id] if i_id < len(row) else None
            if gid_raw is None or str(gid_raw).strip() == "":
                continue
            if isinstance(gid_raw, float) and gid_raw == int(gid_raw):
                case_id = str(int(gid_raw))
            else:
                case_id = str(gid_raw).strip()

            if category not in buckets:
                buckets[category] = []
                order.append(category)

            user_msgs = _split_user_messages(row[i_msg] if i_msg < len(row) else None)
            exp = row[i_exp] if i_exp < len(row) else ""
            expected_outcome = str(exp).strip() if exp is not None else ""
            tool_flag = _parse_tool_flag(row[i_tool] if i_tool is not None and i_tool < len(row) else None)
            if tool_flag is None:
                tool_flag = False
            tparam = row[i_tool_arg] if i_tool_arg is not None and i_tool_arg < len(row) else None
            tool_parameters = str(tparam).strip() if tparam is not None and str(tparam).strip() else None
            tname_cell = row[i_tool_name] if i_tool_name is not None and i_tool_name < len(row) else None
            tool_name = str(tname_cell).strip() if tname_cell is not None and str(tname_cell).strip() else None

            buckets[category].append(
                DialogueTurn(
                    sheet_row=row_idx,
                    case_id=case_id,
                    user_messages=user_msgs,
                    expected_outcome=expected_outcome,
                    expect_tool_call=tool_flag,
                    tool_name=tool_name,
                    tool_parameters=tool_parameters,
                )
            )

        return [
            DialogueGroup(group_id=category, category=category, turns=buckets[category])
            for category in order
        ]
    finally:
        wb.close()


@dataclass
class Config:
    """全仓库固定的连接与裁判配置。各玩具只覆盖用例、工具描述、token。"""

    # OTA 激活地址，用来换 WebSocket url/token
    ota_url: str
    # 禁用 OTA 时直连的 WebSocket；OTA 成功后会被服务端返回值覆盖
    ws_url: str
    # 设备 MAC，同时作为 OTA Device-Id 和 hello.device_mac
    device_mac: str
    # 协议字段 client-id
    client_id: str
    # hello 报文中的设备名称
    device_name: str
    # 玩具服务令牌（来自 toys/<id>/toy.json 或 TOY_AGENT_TOKEN）
    token: str
    # 当前玩具的 Excel 用例
    cases_xlsx: Path
    # 当前玩具的 MCP 工具描述
    mcp_tools_file: Path
    # 报告输出目录，默认 reports/<玩具id>/
    reports_dir: Path
    # 单次 WebSocket 接收超时（秒）
    recv_timeout: float
    # 对话组之间的冷却间隔（秒）
    send_interval: float
    # 每次用户发送后采集回复的最长等待（秒）
    response_collect_timeout: float
    # Qwen 裁判模型名
    judge_model: str
    # 双裁判平均分达到该值即通过（默认 0.7）
    judge_pass_threshold: float
    # Qwen/DashScope API Key；空则读 DASHSCOPE_API_KEY
    api_key: str = ""
    # Qwen OpenAI 兼容地址；空则读 DASHSCOPE_BASE_URL 或百炼默认地址
    judge_base_url: str = ""
    # DeepSeek 裁判模型名
    deepseek_judge_model: str = "deepseek-v4-pro"
    # DeepSeek API Key；空则读 DEEPSEEK_API_KEY
    deepseek_api_key: str = ""
    # DeepSeek OpenAI 兼容地址
    deepseek_base_url: str = "https://api.deepseek.com"
    # 是否启用第二裁判 DeepSeek
    enable_deepseek_judge: bool = True
    # 结果 xlsx 路径；空则写到 reports 目录
    export_xlsx: Optional[Path] = None
    # hello 后每次排空等待（秒）
    hello_drain_per_wait: float = 0.25
    # hello 后连续空闲多少轮才结束排空
    hello_drain_idle_rounds: int = 4
    # hello 排空最多轮数
    hello_drain_max_rounds: int = 16
    # hello 后等待 tools/list 的超时（秒）
    hello_mcp_ready_timeout: float = 8.0
    # 采集阶段单次 recv 超时（秒）
    collect_recv_timeout: float = 0.25
    # 采集静默轮数（兼容旧逻辑，当前以秒为主）
    collect_quiet_rounds: int = 8
    # TTS 结束后再静默多久才进入下一题（秒）
    collect_quiet_seconds: float = 2.0
    # 采集第二阶段上限（秒，fast 模式会缩短）
    collect_phase2_max: float = 2.5
    # 采集收尾余量（秒）
    collect_slack_seconds: float = 0.2
    # 未收到 tools/call 时，tts stop 后再等多久（秒）
    mcp_no_call_grace_seconds: float = 0.8
    # 收到工具结果后静默多久视为本轮结束（秒）
    mcp_completion_quiet_seconds: float = 0.8
    # tts stop 后再排空多久（秒）
    mcp_post_stop_drain_seconds: float = 0.25
    # 第二阶段 LLM 评分并发数
    judge_concurrency: int = 6
    # OTA 请求体 board.ip
    board_ip: str = "192.168.77.46"
    # OTA 请求体 board.type
    board_type: str = "bread-compact-wifi"
    # OTA 请求体 version
    ota_version: int = 0
    # OTA 请求体 application.version
    firmware_version: str = DEFAULT_FIRMWARE_VERSION
    # OTA 请求体 application.name
    application_name: str = "xiaozhi-web-test"
    # 是否使用随机 MAC（当前默认关闭）
    device_mac_random: bool = False
    # 当前玩具 id，对应 toys/<id>/
    toy_id: str = DEFAULT_TOY_ID
    # 可选范围；默认空表示全量，保持原 CLI 行为。
    selected_categories: List[str] = field(default_factory=list)
    selected_case_ids: List[str] = field(default_factory=list)
    selected_case_rows: List[int] = field(default_factory=list)
    judge_template_file: Path = field(default_factory=lambda: Path(__file__).resolve().parent / "templates/template.md")
    thresholds_file: Path = field(default_factory=lambda: Path(__file__).resolve().parent / "settings/evaluation_thresholds.json")
    rules_snapshot: Dict[str, Any] = field(default_factory=dict)
    judge_category: str = "基础问答"
    trigger_type: str = "manual_cli"
    run_id: str = ""
    tools_source: Dict[str, Any] = field(default_factory=dict)
    case_snapshot: Optional[Path] = None
    environment: str = ""
    selected_dataset_ids: List[str] = field(default_factory=list)
    dataset_snapshot: Dict[str, Any] = field(default_factory=dict)
    toy_snapshot: Dict[str, Any] = field(default_factory=dict)
    selected_judge_ids: Optional[List[str]] = None
    judge_endpoints_snapshot: List[Dict[str, Any]] = field(default_factory=list)
    judge_credentials: Dict[str, str] = field(default_factory=dict)

# 连接与消息搜集
@dataclass
class RuntimeState:
    websocket: Any = None
    session_id: str = ""
    response_complete_event: asyncio.Event = field(default_factory=asyncio.Event)
    router_logs: List[Dict[str, Any]] = field(default_factory=list)
    mcp_tools: List[Dict[str, Any]] = field(default_factory=list) # mcp类别响应结果
    mcp_stats: Dict[str, Any] = field(
        default_factory=lambda: {
            "initialize_requests": 0,
            "tools_list_requests": 0,
            "last_tools_list_reply_count": 0,
        }
    )
    # 服务端「助手侧」可见文本片段，按到达顺序拼接（llm/tts 等均在各自 handler 中写入）
    response_text_segments: List[str] = field(default_factory=list)

    def reset_response_collection(self) -> None:
        self.response_text_segments.clear()
        self.response_complete_event.clear()

    def get_full_server_response_text(self) -> str:
        return "".join(self.response_text_segments)


def dialogue_group_to_dict(group: DialogueGroup) -> Dict[str, Any]:
    """将对话组序列化，便于写入报告。"""
    return {
        "group_id": group.group_id,
        "category": group.category,
        "turns": [
            {
                "sheet_row": t.sheet_row,
                "case_id": t.case_id,
                "user_messages": t.user_messages,
                "expected_outcome": t.expected_outcome,
                "expect_tool_call": t.expect_tool_call,
                "tool_name": t.tool_name,
                "tool_parameters": t.tool_parameters,
            }
            for t in group.turns
        ],
    }

# 测试数据基本信息统计
def summarize_dialogue_groups(groups: List[DialogueGroup]) -> Dict[str, Any]:
    by_category: Dict[str, int] = {}
    empty_message_rows = 0
    total_turns = 0
    for g in groups:
        by_category[g.category] = by_category.get(g.category, 0) + len(g.turns)
        for t in g.turns:
            total_turns += 1
            if len(t.user_messages) == 0:
                empty_message_rows += 1
    return {
        "group_count": len(groups),
        "turn_count": total_turns,
        "empty_message_rows": empty_message_rows,
        "by_category": by_category,
    }


@dataclass
class TestContext:
    """单次测试运行期共享上下文（协议、状态、流水线）。"""

    cfg: Config
    state: RuntimeState
    transport: Any
    pipeline: Any
