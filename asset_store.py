"""带内容冲突检查、备份和原子替换的本地资产写入。"""
from __future__ import annotations

import hashlib
import json
import os
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path
from uuid import uuid4

import openpyxl

from config import _dialogue_worksheet, load_dialogue_groups_from_xlsx, _parse_tool_flag
from category_store import list_case_types

ROOT = Path(__file__).resolve().parent


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def atomic_update(path, expected_hash, builder, allowed_root=ROOT):
    path = Path(path).resolve()
    if not path.is_relative_to(Path(allowed_root).resolve()) or not path.is_file():
        raise ValueError("只允许修改项目内已有配置或用例文件")
    lock = path.with_name(path.name + ".edit.lock")
    temporary = path.with_name("." + uuid4().hex + path.suffix)
    try:
        descriptor = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    except FileExistsError:
        raise ValueError("该文件正在保存，请稍后重试") from None
    os.close(descriptor)
    try:
        if digest(path) != expected_hash:
            raise ValueError("文件已被其他操作修改，请重新加载后再保存")
        builder(path, temporary)
        temporary.chmod(path.stat().st_mode & 0o777)
        backup_dir = path.parent / ".backups"
        if not backup_dir.resolve().is_relative_to(Path(allowed_root).resolve()):
            raise ValueError("备份目录不在允许的项目范围内")
        backup_dir.mkdir(exist_ok=True, mode=0o700)
        backup = backup_dir / (datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + uuid4().hex[:8] + "_" + path.name)
        shutil.copy2(path, backup)
        from contextlib import nullcontext
        from dataset_thresholds import case_threshold_transaction
        transaction = case_threshold_transaction(path, temporary, allowed_root) if path.suffix.lower() == ".xlsx" else nullcontext()
        replaced = False
        try:
            with transaction:
                temporary.replace(path)
                replaced = True
        except Exception:
            if replaced:
                shutil.copy2(backup, temporary)
                temporary.replace(path)
            raise
        return backup
    finally:
        temporary.unlink(missing_ok=True)
        lock.unlink(missing_ok=True)


def save_json(path, value, expected_hash, validator=None, allowed_root=ROOT):
    def build(original, temporary):
        temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        if validator:
            validator(temporary)
    return atomic_update(path, expected_hash, build, allowed_root)


def editable_cases(path):
    # 先使用权威解析器验证表头；编辑视图保留原表中的无编号/空字段行，允许修复而非静默丢弃。
    load_dialogue_groups_from_xlsx(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = _dialogue_worksheet(workbook)
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, ())
        columns = {str(v).strip(): i for i, v in enumerate(header) if v is not None}
        def value(row, *names):
            for name in names:
                if name in columns and columns[name] < len(row):
                    return row[columns[name]]
            return None
        def text(row, *names):
            raw = value(row, *names)
            return "" if raw is None else str(raw)
        output, last_category = [], ""
        for number, row in enumerate(rows, 2):
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            category = value(row, "类型")
            if category is not None and str(category).strip():
                last_category = str(category).strip()
            identity = value(row, "用例编号")
            if isinstance(identity, float) and identity.is_integer():
                identity = int(identity)
            output.append({"源行号": number, "类型": last_category, "用例编号": str(identity).strip() if identity is not None else "",
                           "消息列表": text(row, "消息列表"), "预期结果": text(row, "预期结果", "预期结果/非预期结果"),
                           "是否工具调用": bool(_parse_tool_flag(value(row, "是否工具调用"))),
                           "工具名称": text(row, "工具名称"),
                           "工具参数文字描述": text(row, "工具参数文字描述", "工具参数")})
        return output
    finally:
        workbook.close()


def case_issues(path, root=None):
    categories = {entry["name"] for entry in list_case_types(root)}
    issues, seen = [], set()
    for row in editable_cases(path):
        prefix = f'第 {row["源行号"]} 行'
        for key in ["类型", "用例编号", "消息列表", "预期结果"]:
            if not row[key].strip():
                issues.append(f"{prefix}：缺少{key}")
        if row["类型"] and row["类型"] not in categories:
            issues.append(f'{prefix}：未知类型 {row["类型"]}')
        identity = (row["类型"], row["用例编号"])
        if row["用例编号"] and identity in seen:
            issues.append(f"{prefix}：同类型编号重复")
        seen.add(identity)
    return issues


def save_cases(path, rows, expected_hash, allowed_root=ROOT, *, root=None):
    categories = {entry["name"] for entry in list_case_types(root if root is not None else allowed_root)}
    if not rows:
        raise ValueError("不能保存空用例集")
    identities, source_rows = set(), set()
    for row in rows:
        for key in ("类型", "用例编号", "消息列表", "预期结果"):
            if not str(row.get(key) or "").strip():
                raise ValueError(f"用例缺少必填字段：{key}")
        if row["类型"] not in categories:
            raise ValueError("用例类型不受支持")
        identity = (row["类型"], str(row["用例编号"]))
        if identity in identities:
            raise ValueError("同类型的用例编号不能重复")
        identities.add(identity)
        if row.get("源行号") is not None:
            original_row = int(row["源行号"])
            if original_row in source_rows:
                raise ValueError("源行号重复，请重新加载")
            source_rows.add(original_row)

    def build(original, temporary):
        existing = {r["源行号"] for r in editable_cases(original)}
        if source_rows - existing:
            raise ValueError("源行号不存在，请重新加载")
        workbook = openpyxl.load_workbook(original)
        try:
            sheet = _dialogue_worksheet(workbook)
            headers = {str(c.value).strip(): c.column for c in sheet[1] if c.value is not None}
            alias = {"预期结果": "预期结果" if "预期结果" in headers else "预期结果/非预期结果",
                     "工具参数文字描述": "工具参数文字描述" if "工具参数文字描述" in headers else "工具参数" if "工具参数" in headers else "工具参数文字描述"}
            type_column = headers["类型"]
            for merged in list(sheet.merged_cells.ranges):
                if merged.min_col == merged.max_col == type_column and merged.min_row >= 2:
                    style = copy(sheet.cell(merged.min_row, type_column)._style)
                    first, last = merged.min_row, merged.max_row
                    sheet.unmerge_cells(str(merged))
                    for row_number in range(first, last + 1):
                        sheet.cell(row_number, type_column)._style = copy(style)
            for key in ["工具名称", "是否工具调用", alias["工具参数文字描述"]]:
                if key not in headers:
                    headers[key] = sheet.max_column + 1
                    sheet.cell(1, headers[key], key)
            # 先更新现有行，其他列及样式保持；删除后清晰重写类别以消除继承歧义。
            for row in rows:
                if row.get("源行号") is None:
                    continue
                for key in ["类型", "用例编号", "消息列表", "预期结果", "是否工具调用", "工具名称", "工具参数文字描述"]:
                    value = row.get(key)
                    if key == "是否工具调用":
                        value = "是" if value else "否"
                    sheet.cell(int(row["源行号"]), headers[alias.get(key, key)]).value = value
            for row_number in sorted(existing - source_rows, reverse=True):
                sheet.delete_rows(row_number)
            for row in rows:
                if row.get("源行号") is not None:
                    continue
                number = sheet.max_row + 1
                for key in ["类型", "用例编号", "消息列表", "预期结果", "是否工具调用", "工具名称", "工具参数文字描述"]:
                    value = row.get(key)
                    sheet.cell(number, headers[alias.get(key, key)]).value = ("是" if value else "否") if key == "是否工具调用" else value
            start, last_category = 2, sheet.cell(2, type_column).value
            for row_number in range(3, sheet.max_row + 2):
                category = sheet.cell(row_number, type_column).value if row_number <= sheet.max_row else None
                if category != last_category:
                    if last_category and row_number - start > 1:
                        sheet.merge_cells(start_row=start, end_row=row_number - 1, start_column=type_column, end_column=type_column)
                    start, last_category = row_number, category
            workbook.save(temporary)
        finally:
            workbook.close()
        loaded = load_dialogue_groups_from_xlsx(temporary)
        if sum(len(g.turns) for g in loaded) != len(rows):
            raise ValueError("保存后校验失败，未替换原文件")

    return atomic_update(path, expected_hash, build, allowed_root)


def restore_cases(path, backup, expected_hash, allowed_root=ROOT):
    path, backup = Path(path).resolve(), Path(backup).resolve()
    if backup.parent != path.parent / ".backups" or not backup.name.endswith(path.name):
        raise ValueError("备份不属于当前用例文件")
    def build(original, temporary):
        shutil.copy2(backup, temporary)
        load_dialogue_groups_from_xlsx(temporary)
    return atomic_update(path, expected_hash, build, allowed_root)
