"""Management and snapshot regressions: all data stays in temporary directories."""
import json
import os
import tempfile
import unittest
from dataclasses import dataclass, field
from pathlib import Path
from unittest.mock import patch

import openpyxl

from dataset_store import (HEADERS, create_dataset, list_datasets, load_registry,
    prepare_dataset_snapshot, registry_hash, resolve_datasets, save_registry)
from toy_management import create_toy, file_hash, update_toy
from toy_profiles import apply_toy_environment, load_toy_profile


@dataclass
class TestConfig:
    cases_xlsx: Path
    toy_id: str = "default"
    token: str = ""
    environment: str = ""
    ota_url: str = "https://old.invalid/ota"
    device_name: str = "old"
    case_snapshot: object = None
    dataset_snapshot: dict = field(default_factory=dict)
    toy_snapshot: dict = field(default_factory=dict)
    selected_dataset_ids: list = field(default_factory=list)


class ManagementTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        (self.root / "data").mkdir()
        self.source = self.root / "data" / "测试样例.xlsx"
        self.write_cases(self.source)
        self.toy = create_toy("Default", toy_id="default", root=self.root)

    def write_cases(self, path, rows=None):
        wb = openpyxl.Workbook()
        wb.active.title = "Test case"
        wb.active.append(HEADERS)
        for row in rows if rows is not None else [("基础问答", "1", "hello", "answer", "否")]:
            wb.active.append(row)
        wb.save(path)
        wb.close()

    def test_legacy_load_and_token_fallback(self):
        path = self.toy.root / "toy.json"
        path.write_text(json.dumps({"cases_xlsx": "../../data/测试样例.xlsx"}), encoding="utf-8")
        toy = load_toy_profile("default", self.root / "toys")
        self.assertTrue(toy.enabled)
        self.assertFalse(toy.configured)
        with patch.dict(os.environ, {"TOY_AGENT_TOKEN": "test-only"}, clear=True):
            self.assertEqual(toy.resolve_token(), "test-only")
        self.assertEqual(len(list_datasets(toy)), 1)
        self.assertFalse((self.root / "settings").exists())

    def test_reference_does_not_fall_back(self):
        toy = create_toy("Ref", credential_ref="MISSING_TEST_ENV", root=self.root)
        with patch.dict(os.environ, {"TOY_AGENT_TOKEN": "test-only"}, clear=True):
            with self.assertRaises(ValueError):
                toy.resolve_token()
        self.assertEqual(toy.resolve_token("explicit-test-input"), "explicit-test-input")
        self.assertNotIn("token", json.loads((toy.root / "toy.json").read_text(encoding="utf-8")))

    def test_environment_and_capabilities_replace(self):
        toy = create_toy("Environment", capabilities={"基础问答": True, "MCP调用": "unverified"},
            environments={"qa": {"device_name": "qa-device", "credential_ref": "QA_TEST_REF"}}, root=self.root)
        cfg = TestConfig(self.source)
        with patch.dict(os.environ, {"QA_TEST_REF": "test-only"}, clear=True):
            updated = apply_toy_environment(cfg, toy, "qa")
        self.assertEqual(cfg.device_name, "old")
        self.assertEqual(updated.device_name, "qa-device")
        self.assertEqual(updated.toy_snapshot["capabilities"]["基础问答"], "supported")
        self.assertTrue(updated.toy_snapshot["configured"])
        with self.assertRaises(ValueError):
            apply_toy_environment(cfg, toy, "missing")

    def test_validation_disable_backup_conflict(self):
        path = self.toy.root / "toy.json"
        original_hash = file_hash(path)
        updated = update_toy(self.toy, expected_hash=original_hash, enabled=False)
        self.assertFalse(updated.enabled)
        self.assertTrue(list((path.parent / ".backups").iterdir()))
        with self.assertRaises(ValueError):
            update_toy(updated, expected_hash=original_hash, name="stale")
        with self.assertRaises(ValueError):
            apply_toy_environment(TestConfig(self.source), updated)
        with self.assertRaises(ValueError):
            create_toy("invalid", toy_id="../escape", root=self.root)
        with self.assertRaises(ValueError):
            create_toy("invalid", credential_ref="raw key value", root=self.root)
        with self.assertRaises(ValueError):
            create_toy("invalid", environments={"qa": {"ota_url": "https://u:p@example.invalid"}}, root=self.root)
        with self.assertRaises((ValueError, FileExistsError)):
            create_toy("duplicate", toy_id="default", root=self.root)

    def test_copy_registry_scope_and_conflicts(self):
        public = create_dataset("Public", source_path=self.source, root=self.root)
        private = create_dataset("Private", scope="toy", owner_toy_id="default", root=self.root)
        other = create_toy("Other", root=self.root)
        other_set = create_dataset("Other only", scope="toy", owner_toy_id=other.id, root=self.root)
        self.assertEqual(public.path.read_bytes(), self.source.read_bytes())
        available = {d.id for d in list_datasets(self.toy)}
        self.assertIn(private.id, available)
        self.assertNotIn(other_set.id, available)
        with self.assertRaises(ValueError):
            resolve_datasets(self.toy, [other_set.id])
        old_hash = registry_hash(self.root)
        entries = load_registry(self.root)
        entries[0].enabled = False
        save_registry(entries, old_hash, self.root)
        with self.assertRaises(ValueError):
            save_registry(entries, old_hash, self.root)
        with self.assertRaises(ValueError):
            resolve_datasets(self.toy, [public.id])
        with self.assertRaises(ValueError):
            save_registry(entries + entries, registry_hash(self.root), self.root)

    def test_single_snapshot_and_empty_rejection(self):
        cfg = TestConfig(self.source)
        updated = prepare_dataset_snapshot(cfg, self.root / "run", root=self.root)
        self.assertEqual(updated.cases_xlsx.read_bytes(), self.source.read_bytes())
        self.assertEqual(updated.dataset_snapshot["rows"]["2"]["row"], 2)
        self.assertEqual(updated.dataset_snapshot["sessions"][0]["rows"], [2])
        empty = create_dataset("Empty", root=self.root)
        cfg.selected_dataset_ids = [empty.id]
        with self.assertRaisesRegex(ValueError, "空测试集"):
            prepare_dataset_snapshot(cfg, self.root / "empty-run", root=self.root)

    def test_merge_source_mapping_and_sessions(self):
        first = create_dataset("A", source_path=self.source, root=self.root)
        second = create_dataset("B", source_path=self.source, root=self.root)
        cfg = TestConfig(self.source, selected_dataset_ids=[first.id, second.id])
        updated = prepare_dataset_snapshot(cfg, self.root / "run", root=self.root)
        snapshot = updated.dataset_snapshot
        self.assertEqual(len(snapshot["sessions"]), 2)
        self.assertEqual(snapshot["rows"]["2"]["dataset_id"], first.id)
        self.assertEqual(snapshot["rows"]["3"]["dataset_id"], second.id)
        self.assertEqual(snapshot["rows"]["3"]["row"], 2)
        wb = openpyxl.load_workbook(updated.cases_xlsx)
        self.assertEqual(wb.active.max_row, 3)
        self.assertEqual(wb.active.cell(3, 1).value, "基础问答")
        wb.close()

    def test_reject_duplicate_source_and_path_escape(self):
        self.write_cases(self.source, [("基础问答", "1", "a", "b"), ("基础问答", "1", "c", "d")])
        with self.assertRaisesRegex(ValueError, "重复"):
            prepare_dataset_snapshot(TestConfig(self.source), self.root / "run", self.root)
        path = self.toy.root / "toy.json"
        path.write_text(json.dumps({"cases_xlsx": "../../../outside.xlsx"}), encoding="utf-8")
        with self.assertRaises(ValueError):
            load_toy_profile("default", self.root / "toys")
        with self.assertRaises(ValueError):
            create_dataset("outside", source_path="../outside.xlsx", root=self.root)


if __name__ == "__main__":
    unittest.main()
