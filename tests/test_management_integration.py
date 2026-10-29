"""跨测试集/工具/报告/UI 的临时隔离回归；所有外部传输均 mock。"""
import asyncio
import json
import tempfile
import unittest
from contextlib import ExitStack
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import openpyxl
import dataset_store
from config import Config
from dataset_store import HEADERS, create_dataset, prepare_dataset_snapshot
from text_input_test import run_framework
from toy_management import create_toy
from toy_profiles import apply_toy_environment
from ui_request import validate_request
from ui_results import read_report

PROJECT = Path(__file__).resolve().parents[1]
TOOLS = [{"name": "weather", "description": "模拟天气定义", "inputSchema": {"type": "object"}}]
RULES = {"thresholds": {}, "case_thresholds": [], "rubrics": {}, "overrides": {}}


class ManagementIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="management-integration-")
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        (self.root / "data").mkdir()
        self.cases = self.root / "data" / "测试样例.xlsx"
        self.write_cases(self.cases)
        self.tools = self.root / "data" / "工具描述.txt"
        self.tools.write_text(json.dumps(TOOLS), encoding="utf-8")
        self.toy = create_toy("隔离测试玩具", toy_id="integration", root=self.root,
                              capabilities={"text": True, "tool_injection": True, "tool_observation": True})
        self.cfg = Config("http://127.0.0.1/mock-ota", "ws://127.0.0.1/mock-ws", "02:00:00:00:00:01", "mock", "mock", "mock-token",
                          self.cases, self.tools, self.root / "reports", .1, 0, .1, "mock-model", .7,
                          toy_id=self.toy.id, rules_snapshot=RULES, tools_source={"type": "local"},
                          export_xlsx=self.root / "reports" / "output.xlsx")
        self.cfg = apply_toy_environment(self.cfg, self.toy)
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.stack.enter_context(patch.object(dataset_store, "PROJECT_ROOT", self.root))

    @staticmethod
    def write_cases(path, name="weather"):
        book = openpyxl.Workbook()
        book.active.title = "Test case"
        book.active.append(HEADERS)
        book.active.append(("MCP调用", "same-id", "天气如何", "调用天气工具", "是", name, "仅供展示"))
        book.save(path)
        book.close()

    def datasets(self):
        public = create_dataset("公共集", source_path=self.cases, dataset_id="shared", root=self.root)
        owned = create_dataset("专属集", "toy", self.toy.id, source_path=self.cases, dataset_id="owned", root=self.root)
        return [public, owned]

    def test_pure_rule_request_without_keys_and_invalid_expected_blocks(self):
        self.assertFalse(self.cfg.api_key)
        self.assertFalse(self.cfg.deepseek_api_key)
        groups, _ = validate_request(self.cfg)
        self.assertEqual(len(groups), 1)
        self.write_cases(self.cases, "not-declared")
        with self.assertRaisesRegex(ValueError, "未在|不存在|定义"):
            validate_request(self.cfg)

    def test_unsupported_capability_blocks(self):
        unsupported = replace(self.cfg, toy_snapshot={"configured": True, "capabilities": {"text": "supported"}})
        with self.assertRaisesRegex(ValueError, "能力不支持|未验证"):
            validate_request(unsupported)

    def run_mock(self, cfg, failed_reply=False):
        sockets = []
        class Transport:
            def __init__(self, cfg, state, router):
                self.state = state
            async def send_ota_request(self):
                return {"ok": True, "status": 200, "data": {"websocket": {"url": "ws://127.0.0.1/mock"}}}
            @staticmethod
            def extract_websocket_from_ota(data):
                return data["websocket"]["url"], ""
            async def connect_websocket(self, *args):
                self.state.websocket = SimpleNamespace(close=AsyncMock())
                sockets.append(self.state.websocket)
                return True
            async def send_hello_message(self):
                return True
        async def collect(inspector, start):
            return {"saw_mcp": True, "mcp_calls": [{"tool_name": "weather", "tool_args": {}}],
                    "mcp_reply_confirmed": not failed_reply, "response_complete": True}
        with patch("text_input_test.TransportClient", Transport), patch("collection.TextInteractionPipeline.send_text_message", new=AsyncMock()), \
             patch("mcp.MCPInspector.collect_result_after_send", new=collect), \
             patch("judge.DualLLMJudge", side_effect=AssertionError("纯规则不应初始化裁判")):
            output = asyncio.run(run_framework(cfg))
        for socket in sockets:
            socket.close.assert_awaited_once()
        return output, sockets

    def test_two_same_category_and_id_datasets_isolate_sessions_and_export_sources(self):
        datasets = self.datasets()
        snapshot = prepare_dataset_snapshot(replace(self.cfg, selected_dataset_ids=[d.id for d in datasets]), self.root / "snapshot")
        groups, _ = validate_request(snapshot)
        self.assertEqual(len(groups), 2)
        self.assertEqual([g.turns[0].case_id for g in groups], ["same-id", "same-id"])
        output, sockets = self.run_mock(snapshot)
        self.assertEqual(output["execution_status"], "completed")
        self.assertEqual(len(sockets), 2)
        self.assertIsNot(sockets[0], sockets[1])
        report = read_report(Path(output["report_paths"]["detail_report"]), self.root)
        self.assertEqual(report["metrics"]["case_count"], 2)
        self.assertEqual(report["metrics"]["passed_count"], 2)
        self.assertEqual(report["metrics"]["scored_count"], 0)
        self.assertEqual(report["metrics"]["pass_rate"], 1)
        self.assertEqual({row["source"]["dataset_id"] for row in report["rows"]}, {"shared", "owned"})
        self.assertEqual({row["source"]["row"] for row in report["rows"]}, {2})
        book = openpyxl.load_workbook(snapshot.export_xlsx)
        self.addCleanup(book.close)
        rows = list(book.active.values)
        header = {value: index for index, value in enumerate(rows[0])}
        self.assertEqual({row[header["测试集"]] for row in rows[1:]}, {"公共集", "专属集"})
        self.assertEqual([row[header["源行号"]] for row in rows[1:]], [2, 2])
        self.assertEqual([row[header["规则得分"]] for row in rows[1:]], [1, 1])
        self.assertEqual([row[header["测试结果是否通过"]] for row in rows[1:]], ["通过", "通过"])

    def test_tool_reply_execution_error_excluded_from_business_denominator(self):
        output, _ = self.run_mock(self.cfg, failed_reply=True)
        self.assertEqual(output["execution_status"], "failed")
        report = read_report(Path(output["report_paths"]["detail_report"]), self.root)
        self.assertEqual(report["rows"][0]["status"], "未执行")
        self.assertEqual(report["rows"][0]["rule_score"], 1)
        self.assertTrue(report["rows"][0]["execution_error"])
        self.assertEqual(report["metrics"]["evaluated_count"], 0)
        self.assertEqual(report["metrics"]["unevaluated_count"], 1)
        self.assertIsNone(report["metrics"]["pass_rate"])
        self.assertEqual(report["metrics"]["failed_count"], 0)

    def test_dataset_sidecar_threshold_drives_run_rows_and_legacy_fallback(self):
        from dataset_thresholds import save_dataset_thresholds, snapshot_threshold, thresholds_path
        from toy_management import file_hash
        base = self.root / "data" / "base.xlsx"
        book = openpyxl.Workbook()
        book.active.title = "Test case"
        book.active.append(HEADERS)
        book.active.append(("基础问答", "plain", "你好", "回复问候", "否", "", ""))
        book.save(base)
        book.close()
        dataset = create_dataset("阈值集", source_path=base, dataset_id="thresholds", root=self.root)
        # Sidecar is authoritative: 0.5 must beat the 0.7 category default.
        save_dataset_thresholds(dataset, {"case_thresholds": [
            {"dataset_id": dataset.id, "category": "基础问答", "case_id": "plain",
             "sheet_row": 2, "threshold": .5, "expect_tool_call": False}]},
            file_hash(thresholds_path(dataset, self.root)), root=self.root)
        cfg = prepare_dataset_snapshot(replace(self.cfg, selected_dataset_ids=[dataset.id]), self.root / "run")
        self.assertEqual(snapshot_threshold(cfg.dataset_snapshot, "基础问答", "plain", 2, root=self.root, cfg=cfg), .5)
        # Without a sidecar the page keeps the legacy category threshold.
        plain = self.datasets()[0]
        fallback = prepare_dataset_snapshot(replace(self.cfg, selected_dataset_ids=[plain.id]), self.root / "run2")
        self.assertIsNone(snapshot_threshold(fallback.dataset_snapshot, "MCP调用", "same-id", 2, root=self.root, cfg=fallback))
        # Missing source mapping stays None instead of guessing.
        self.assertIsNone(snapshot_threshold({}, "基础问答", "plain", 2, root=self.root))

    def test_scoped_rules_reject_ambiguous_and_project_dataset_original_rows(self):
        from evaluation_selection import scoped_rules
        datasets = self.datasets()
        cfg = prepare_dataset_snapshot(replace(self.cfg, selected_dataset_ids=[d.id for d in datasets]), self.root / "snapshot")
        rule = {"toy_id": self.toy.id, "category": "MCP调用", "case_id": "same-id", "sheet_row": 2, "threshold": .81}
        with self.assertRaisesRegex(ValueError, "歧义"):
            scoped_rules(cfg, dict(RULES, case_thresholds=[rule]))
        result = scoped_rules(cfg, dict(RULES, case_thresholds=[dict(rule, dataset_id="owned")]))
        self.assertEqual(len(result["case_thresholds"]), 1)
        projected = result["case_thresholds"][0]
        self.assertEqual(projected["sheet_row"], 3)
        self.assertEqual(projected["dataset_id"], "owned")
        self.assertEqual(projected["threshold"], .81)
        self.assertEqual(cfg.dataset_snapshot["rows"]["3"]["row"], 2)

    def test_apptest_real_toy_and_dataset_saves_in_temporary_project(self):
        from streamlit.testing.v1 import AppTest
        import toy_profiles
        import toy_management
        import ui_catalogs
        original_create = dataset_store.create_dataset
        def temporary_create(*args, **kwargs):
            return original_create(*args, **dict(kwargs, root=self.root))
        script = (
            "from pathlib import Path\n"
            "from toy_profiles import load_toy_profile\n"
            "from ui_catalogs import toy_management_page, dataset_management\n"
            f"toy = load_toy_profile('integration', Path({str(self.root / 'toys')!r}))\n"
            "toy_management_page(toy, lambda _: None)\n"
            "dataset_management(toy)\n"
        )
        with patch.object(toy_management, "PROJECT_ROOT", self.root), \
             patch.object(toy_profiles, "PROJECT_ROOT", self.root), \
             patch.object(toy_profiles, "TOYS_DIR", self.root / "toys"), \
             patch.object(ui_catalogs, "create_dataset", side_effect=temporary_create):
            app = AppTest.from_string(script, default_timeout=30).run()
            self.assertEqual(len(app.exception), 0)
            self.assertFalse(any(w.label == "显示名称" for w in app.text_input))
            next(w for w in app.button if w.label == "＋ 接入玩具").click().run()
            self.assertEqual({w.label for w in app.text_input if w.label in {"玩具名称", "玩具 Token"}}, {"玩具名称", "玩具 Token"})
            next(w for w in app.text_input if w.label == "玩具名称").set_value("页面创建玩具")
            next(w for w in app.text_input if w.label == "玩具 Token").set_value("1234")
            next(w for w in app.button if w.label == "接入并保存").click().run()
            self.assertEqual(len(app.exception), 0)
            configs = [json.loads(path.read_text(encoding="utf-8")) for path in (self.root / "toys").glob("*/toy.json")]
            self.assertTrue(any(item["name"] == "页面创建玩具" for item in configs), ([e.value for e in app.error], configs))
            next(w for w in app.button if w.key == "toy_detail_integration").click().run()
            before = json.loads((self.toy.root / "toy.json").read_text(encoding="utf-8"))
            self.assertFalse(any(w.label == "接口与环境配置 JSON" for w in app.text_area))
            next(w for w in app.text_input if w.label == "玩具名称").set_value("页面修改名称")
            next(w for w in app.text_input if w.label == "玩具 Token").set_value("test-only-edited-token")
            next(w for w in app.button if w.label == "保存修改").click().run()
            self.assertEqual(len(app.exception), 0)
            stored = json.loads((self.toy.root / "toy.json").read_text(encoding="utf-8"))
            self.assertEqual(stored["name"], "页面修改名称")
            self.assertTrue(stored["credential_ref"].startswith("local:"))
            updated = toy_profiles.load_toy_profile("integration", self.root / "toys")
            self.assertEqual(updated.resolve_token(), "test-only-edited-token")
            self.assertNotIn("test-only-edited-token", (self.toy.root / "toy.json").read_text(encoding="utf-8"))
            self.assertEqual(stored["capabilities"], before["capabilities"])
            self.assertEqual(stored["environments"], before["environments"])
            self.assertTrue(list((self.toy.root / ".backups").glob("*_toy.json")))
            for scope, name in (("public", "页面公共集"), ("toy", "页面专属集")):
                next(w for w in app.text_input if w.label == "测试集名称").set_value(name)
                next(w for w in app.selectbox if w.label == "测试集归属").set_value(scope)
                next(w for w in app.button if w.label == "创建测试集").click().run()
                self.assertEqual(len(app.exception), 0)
            registry = dataset_store.load_registry(self.root)
            self.assertEqual({item.name for item in registry}, {"页面公共集", "页面专属集"})
            self.assertEqual({item.scope for item in registry}, {"public", "toy"})
            for item in registry:
                self.assertTrue(item.path.is_file())
                if item.scope == "toy":
                    self.assertEqual(item.owner_toy_id, self.toy.id)
                book = openpyxl.load_workbook(item.path)
                try:
                    self.assertEqual(tuple(cell.value for cell in book.active[1]), HEADERS)
                finally:
                    book.close()
            self.assertEqual(len(app.error), 0)

    def test_apptest_tool_import_save_and_backup_real_action(self):
        from streamlit.testing.v1 import AppTest
        import asset_store
        import toy_profiles
        old_text = self.tools.read_text(encoding="utf-8")
        updated = [{"name": "battery", "description": "模拟电池", "inputSchema": {"type": "object", "properties": {}}}]
        original_atomic = asset_store.atomic_update
        def temporary_atomic(path, expected_hash, builder, **kwargs):
            return original_atomic(path, expected_hash, builder, allowed_root=self.root)
        script = (
            "from pathlib import Path\n"
            "from toy_profiles import load_toy_profile\n"
            "from ui_management import tools_management\n"
            f"toy = load_toy_profile('integration', Path({str(self.root / 'toys')!r}))\n"
            "tools_management(toy)\n"
        )
        with patch.object(asset_store, "atomic_update", side_effect=temporary_atomic), \
             patch.object(toy_profiles, "TOYS_DIR", self.root / "toys"):
            app = AppTest.from_string(script, default_timeout=30).run()
            self.assertEqual(len(app.exception), 0)
            next(w for w in app.text_area if w.label == "或粘贴嵌入式工具定义").set_value(json.dumps(updated))
            next(w for w in app.button if w.label == "校验并预览工具定义").click().run()
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            next(w for w in app.button if w.label == "保存嵌入式工具描述").click().run()
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertEqual(json.loads(self.tools.read_text(encoding="utf-8")), updated)
            backups = list((self.tools.parent / ".backups").glob("*_工具描述.txt"))
            self.assertEqual(len(backups), 1)
            self.assertEqual(backups[0].read_text(encoding="utf-8"), old_text)

    def test_app_menu_and_dataset_selection_without_real_writes(self):
        from streamlit.testing.v1 import AppTest
        import toy_profiles
        import toy_management
        import ui_jobs
        self.datasets()
        # Patch module roots before AppTest imports its application namespace.
        with patch.object(toy_profiles, "TOYS_DIR", self.root / "toys"), \
             patch.object(toy_profiles, "PROJECT_ROOT", self.root), \
             patch.object(toy_management, "PROJECT_ROOT", self.root), \
             patch.object(ui_jobs, "ROOT", self.root), \
             patch.object(ui_jobs.JobManager, "busy", return_value=False), \
             patch.object(ui_jobs.JobManager, "refresh", return_value=None), \
             patch.object(ui_jobs.JobManager, "start", side_effect=AssertionError("UI 测试禁止真实启动")):
            app = AppTest.from_file(str(PROJECT / "streamlit_app.py"), default_timeout=30).run()
            self.assertEqual(len(app.exception), 0)
            menu = next(widget for widget in app.radio if widget.label == "工作空间")
            self.assertIn("玩具管理", menu.options)
            menu.set_value("玩具管理").run()
            self.assertEqual(len(app.exception), 0)
            self.assertTrue(any(widget.label == "＋ 接入玩具" for widget in app.button))
            self.assertTrue(any(widget.label == "配置详情 →" for widget in app.button))
            self.assertFalse(any(widget.label == "显示名称" for widget in app.text_input))
            next(widget for widget in app.button if widget.key == "toy_run_integration").click().run()
            self.assertEqual(app.session_state["navigation"], "评测配置与执行")
            self.assertEqual(app.session_state["toy_choice"], "integration")
            self.assertEqual(len(app.exception), 0)
            selector = next(widget for widget in app.multiselect if widget.key == "run_datasets")
            selector.set_value(["shared", "owned"]).run()
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(set(app.session_state["run_datasets"]), {"shared", "owned"})
            self.assertTrue(any("即将评测 2 条" in item.value for item in app.caption))
            self.assertTrue(next(widget for widget in app.button if "开始评测" in widget.label).disabled)


if __name__ == "__main__":
    unittest.main()
