"""补齐功能的实际集成回归，全部使用临时资产与本机测试服务。"""
import asyncio
from dataclasses import replace
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import json
from pathlib import Path
import shutil
import subprocess
import sys
import os
import time
import tempfile
import threading
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from zoneinfo import ZoneInfo

import openpyxl
import pandas as pd

from asset_store import digest, editable_cases, restore_cases, save_cases, case_issues
from config import load_dialogue_groups_from_xlsx
from evaluation_support import case_threshold, load_rules
from execution_lock import device_lock
from judge import LLMJudge, post_judge_results
from test_integration import config, FakeSocket
from tools_source import discover_tools
from trigger_run import execute_event, run_daily_if_due, scan_events, validate_event
from notify_wecom import xlsx_detailed_stats


class CompletionTests(unittest.TestCase):
    def test_invalid_rows_are_visible_and_can_be_repaired(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "invalid.xlsx"
            wb = openpyxl.Workbook()
            wb.active.append(["类型", "用例编号", "消息列表", "预期结果"])
            wb.active.append([None, None, "你好", None])
            wb.save(path)
            wb.close()
            self.assertEqual(len(editable_cases(path)), 1)
            self.assertGreater(len(case_issues(path)), 0)
            row = editable_cases(path)[0]
            row.update({"类型": "基础问答", "用例编号": "fixed", "预期结果": "打招呼"})
            save_cases(path, [row], digest(path), allowed_root=Path(tmp))
            self.assertEqual(case_issues(path), [])
            self.assertEqual(load_dialogue_groups_from_xlsx(path)[0].turns[0].case_id, "fixed")
    def test_running_evaluation_keeps_its_case_snapshot(self):
        from text_input_test import run_framework
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "cases.xlsx"
            shutil.copy2(config().cases_xlsx, path)
            cfg = replace(config(), cases_xlsx=path, reports_dir=Path(tmp) / "reports", selected_case_ids=["1"],
                          response_collect_timeout=0.1, collect_quiet_seconds=0.001, hello_drain_per_wait=0.001,
                          hello_drain_idle_rounds=1, send_interval=0)
            original_hash = digest(path)
            original_expected = editable_cases(path)[0]["预期结果"]
            async def connect(*args, **kwargs):
                rows = editable_cases(path)
                rows[0]["预期结果"] = "下一轮的新预期"
                save_cases(path, rows, digest(path), allowed_root=Path(tmp))
                return FakeSocket()
            response = SimpleNamespace(ok=True, status_code=200, json=lambda: {"websocket": {"url": "ws://test.invalid"}})
            answer = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='{"score":0.9,"reason":"符合","is_off_topic":false}'))])
            with patch("collection.requests.post", return_value=response), patch("collection.websockets.connect", side_effect=connect), \
                 patch.object(LLMJudge, "_create_client", return_value=(object(), None)), patch.object(LLMJudge, "_chat", return_value=answer):
                result = asyncio.run(run_framework(cfg))
            data = json.loads(Path(result["report_paths"]["detail_report"]).read_text())
            self.assertEqual(data["protocol_info"]["cases_sha256"], original_hash)
            self.assertEqual(data["results"][0]["turn_results"][0]["expected_outcome"], original_expected)
            wb = openpyxl.load_workbook(result["report_paths"]["xlsx_export"])
            headers = {c.value: c.column for c in wb.active[1]}
            self.assertEqual(wb.active.cell(2, headers["预期结果"]).value, original_expected)
            wb.close()
            self.assertEqual(editable_cases(path)[0]["预期结果"], "下一轮的新预期")
    def test_remote_tool_source_is_advertised_by_main_engine(self):
        from text_input_test import run_framework
        advertised = []
        class RecordingSocket(FakeSocket):
            async def send(self, raw):
                message = json.loads(raw)
                result = message.get("payload", {}).get("result", {})
                if "tools" in result:
                    advertised.extend(result["tools"])
                await super().send(raw)
        async def connect(*args, **kwargs):
            return RecordingSocket()
        with tempfile.TemporaryDirectory() as tmp:
            cfg = replace(config(), reports_dir=Path(tmp), selected_case_ids=["1"],
                          tools_source={"type": "mcp_http", "url": "http://127.0.0.1:1/mcp"},
                          response_collect_timeout=0.1, collect_quiet_seconds=0.001, send_interval=0,
                          hello_drain_per_wait=0.001, hello_drain_idle_rounds=1)
            tools = [{"name": "remote_only", "description": "远程定义", "inputSchema": {"type": "object"}}]
            response = SimpleNamespace(ok=True, status_code=200, json=lambda: {"websocket": {"url": "ws://test.invalid"}})
            answer = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='{"score":0.9,"reason":"符合","is_off_topic":false}'))])
            with patch("tools_source._remote_tools", return_value=tools), patch("collection.requests.post", return_value=response), \
                 patch("collection.websockets.connect", side_effect=connect), patch.object(LLMJudge, "_create_client", return_value=(object(), None)), \
                 patch.object(LLMJudge, "_chat", return_value=answer):
                result = asyncio.run(run_framework(cfg))
            data = json.loads(Path(result["report_paths"]["detail_report"]).read_text())
            self.assertEqual(advertised, tools)
            self.assertEqual(data["protocol_info"]["tools_origin"]["type"], "mcp_http")
            self.assertEqual(data["results"][0]["mcp_tools_list"]["returned_tool_names"], ["remote_only"])

    def test_daily_run_keeps_report_when_notifier_fails(self):
        import daily_run
        with tempfile.TemporaryDirectory() as tmp:
            calls = []
            def run(command, **kwargs):
                calls.append((command, kwargs))
                if command[1].endswith("text_input_test.py"):
                    target = Path(command[command.index("--export-xlsx") + 1])
                    target.parent.mkdir(parents=True, exist_ok=True)
                    wb = openpyxl.Workbook()
                    wb.active.append(["测试结果是否通过"])
                    wb.active.append(["通过"])
                    wb.save(target)
                    wb.close()
                    return SimpleNamespace(returncode=0)
                self.assertTrue(Path(command[2]).is_file())
                return SimpleNamespace(returncode=2)
            with patch("daily_run.ROOT", Path(tmp)), patch.dict("os.environ", {"TOY_IDS": "default"}), \
                 patch("daily_run.subprocess.run", side_effect=run):
                code = daily_run.main()
            self.assertEqual(code, 2)
            self.assertTrue(list((Path(tmp) / "reports/default").glob("*.xlsx")))
            self.assertEqual(calls[0][1]["env"]["TOY_EVAL_TRIGGER"], "scheduled")
            self.assertNotIn("--api-key", calls[0][0])
            self.assertNotIn("--token", calls[0][0])

    @unittest.skipIf(os.name == "nt", "Windows TerminateProcess 不发送 SIGTERM，待目标 Runner 验证停止路径")
    def test_watcher_real_process_starts_and_stops_without_running_evaluation(self):
        from ui_jobs import ROOT
        with tempfile.TemporaryDirectory() as tmp:
            config_path = Path(tmp) / "config.json"
            config_path.write_text(json.dumps({"daily_enabled": False, "daily_time": "10:00", "timezone": "Asia/Shanghai", "toy_ids": ["default"], "watch_events": True}))
            state_dir = Path(tmp) / "state"
            process = subprocess.Popen([sys.executable, str(ROOT / "trigger_run.py"), "--watch", "--config", str(config_path), "--state-dir", str(state_dir)],
                                       stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
            try:
                deadline = time.monotonic() + 10
                while time.monotonic() < deadline and not (state_dir / "heartbeat.json").exists():
                    if process.poll() is not None:
                        self.fail(process.stderr.read().decode())
                    time.sleep(0.05)
                self.assertEqual(json.loads((state_dir / "heartbeat.json").read_text())["status"], "running")
                self.assertTrue((state_dir / "watch.lock").is_file())
                process.terminate()
                process.wait(timeout=10)
                self.assertEqual(json.loads((state_dir / "heartbeat.json").read_text())["status"], "stopped")
                self.assertFalse((state_dir / "watch.lock").exists())
                self.assertFalse(list((state_dir / "events").glob("*.json")))
            finally:
                if process.poll() is None:
                    process.kill()
                    process.wait()
                process.stderr.close()
    def test_minimal_case_file_optional_tool_columns(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "minimal.xlsx"
            wb = openpyxl.Workbook()
            wb.active.append(["类型", "用例编号", "消息列表", "预期结果"])
            wb.active.append(["基础问答", "A", "你好", "打招呼"])
            wb.active.append(["基础问答", "B", 0, 0])
            wb.save(path)
            wb.close()
            groups = load_dialogue_groups_from_xlsx(path)
            self.assertFalse(groups[0].turns[0].expect_tool_call)
            rows = editable_cases(path)
            rows[0]["消息列表"] = "你好呀"
            save_cases(path, rows, digest(path), allowed_root=Path(tmp))
            self.assertEqual(load_dialogue_groups_from_xlsx(path)[0].turns[0].user_messages, ["你好呀"])
            self.assertEqual(load_dialogue_groups_from_xlsx(path)[0].turns[1].user_messages, ["0"])
            self.assertEqual(load_dialogue_groups_from_xlsx(path)[0].turns[1].expected_outcome, "0")

    def test_ota_failure_is_unexecuted_in_json_excel_and_summary(self):
        from text_input_test import run_framework
        with tempfile.TemporaryDirectory() as tmp:
            cfg = replace(config(), reports_dir=Path(tmp), selected_case_ids=["1"])
            with patch("collection.requests.post", side_effect=ConnectionError("offline")):
                result = asyncio.run(run_framework(cfg))
            self.assertEqual(result["execution_status"], "failed")
            self.assertEqual(xlsx_detailed_stats(Path(result["report_paths"]["xlsx_export"])), (1, 0, 0, 0, 1))
            self.assertIn("未执行 1", Path(result["report_paths"]["summary_report"]).read_text())

    def test_ui_worker_real_subprocess_completes_with_local_services(self):
        from websockets.asyncio.server import serve
        from ui_jobs import JobManager
        loop = asyncio.new_event_loop()
        ready = threading.Event()
        address = {}

        async def socket_handler(socket):
            await socket.recv()
            await socket.send(json.dumps({"type": "hello", "session_id": "local-regression"}))
            await socket.send(json.dumps({"type": "mcp", "payload": {"id": 1, "method": "tools/list"}}))
            while True:
                try:
                    message = json.loads(await socket.recv())
                except Exception:
                    return
                if message.get("type") == "listen":
                    await socket.send(json.dumps({"type": "tts", "state": "stop", "text": "你好呀"}))

        async def host():
            server = await serve(socket_handler, "127.0.0.1", 0)
            address["ws"] = server
            address["port"] = server.sockets[0].getsockname()[1]
            ready.set()

        def thread_main():
            asyncio.set_event_loop(loop)
            loop.run_until_complete(host())
            loop.run_forever()

        ws_thread = threading.Thread(target=thread_main, daemon=True)
        ws_thread.start()
        self.assertTrue(ready.wait(5))

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def do_POST(self):
                self.rfile.read(int(self.headers["Content-Length"]))
                if self.path == "/ota":
                    result = {"websocket": {"url": f'ws://127.0.0.1:{address["port"]}', "token": "local-token"}}
                else:
                    result = {"id": "local", "object": "chat.completion", "created": 0, "model": "local",
                              "choices": [{"index": 0, "finish_reason": "stop", "message": {"role": "assistant", "content": json.dumps({"score": 0.9, "reason": "符合", "is_off_topic": False})}}]}
                content = json.dumps(result).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(content)))
                self.end_headers()
                self.wfile.write(content)

        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        http_thread = threading.Thread(target=server.serve_forever, daemon=True)
        http_thread.start()
        manager = JobManager()
        try:
            with tempfile.TemporaryDirectory() as tmp:
                base = f"http://127.0.0.1:{server.server_port}"
                cfg = replace(config(), reports_dir=Path(tmp), token="local-token", api_key="local-qwen", deepseek_api_key="local-ds",
                              ota_url=base + "/ota", judge_base_url=base + "/v1", deepseek_base_url=base + "/v1",
                              selected_case_ids=["1"], response_collect_timeout=1, collect_quiet_seconds=0.01,
                              hello_drain_per_wait=0.01, hello_drain_idle_rounds=1, send_interval=0,
                              device_mac="local-worker-regression")
                with patch.dict("os.environ", {"NO_PROXY": "127.0.0.1", "no_proxy": "127.0.0.1"}):
                    directory = manager.start(cfg)
                deadline = time.monotonic() + 20
                while time.monotonic() < deadline:
                    manager.refresh()
                    record = json.loads((directory / "job.json").read_text())
                    if record["status"] != "running":
                        break
                    time.sleep(0.05)
                self.assertEqual(record["status"], "completed", record)
                self.assertEqual(xlsx_detailed_stats(Path(record["report_paths"]["xlsx_export"])), (1, 1, 0, 0, 0))
                for process, _ in manager._processes.values():
                    process.wait(timeout=5)
        finally:
            for process, _ in manager._processes.values():
                if process.poll() is None:
                    process.terminate()
                    process.wait(timeout=5)
            server.shutdown()
            server.server_close()
            http_thread.join()
            async def shutdown():
                address["ws"].close()
                await address["ws"].wait_closed()
            asyncio.run_coroutine_threadsafe(shutdown(), loop).result(timeout=5)
            loop.call_soon_threadsafe(loop.stop)
            ws_thread.join()
            loop.close()
    def test_edit_add_delete_restore_and_conflict(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "cases.xlsx"
            shutil.copy2(config().cases_xlsx, path)
            initial = digest(path)
            rows = editable_cases(path)
            rows[0]["预期结果"] = "修改后的预期"
            removed = rows.pop(1)
            rows.append(dict(rows[0], **{"源行号": None, "用例编号": "NEW-1"}))
            backup = save_cases(path, rows, initial, allowed_root=root)
            updated = editable_cases(path)
            self.assertEqual(len(updated), 98)
            self.assertEqual(updated[0]["预期结果"], "修改后的预期")
            self.assertTrue(any(r["用例编号"] == "NEW-1" for r in updated))
            self.assertFalse(any(r["用例编号"] == removed["用例编号"] for r in updated))
            with self.assertRaisesRegex(ValueError, "其他操作"):
                save_cases(path, rows, initial, allowed_root=root)
            restore_cases(path, backup, digest(path), allowed_root=root)
            self.assertEqual(digest(path), initial)

    def test_invalid_edit_does_not_replace_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "cases.xlsx"
            shutil.copy2(config().cases_xlsx, path)
            initial = digest(path)
            rows = editable_cases(path)
            rows[0]["类型"] = "未知分类"
            with self.assertRaises(ValueError):
                save_cases(path, rows, initial, allowed_root=Path(tmp))
            self.assertEqual(digest(path), initial)

    def test_single_case_threshold_changes_final_judgement(self):
        cfg = config()
        rules = load_rules(cfg)
        rules["case_thresholds"] = [{"toy_id": cfg.toy_id, "category": "基础问答", "case_id": "1", "threshold": 0.95},
                                    {"toy_id": cfg.toy_id, "category": "基础问答", "case_id": "1", "sheet_row": 2, "threshold": 0.99}]
        self.assertEqual(case_threshold(rules, cfg.toy_id, "基础问答", "1", 2), 0.99)
        self.assertEqual(case_threshold(rules, cfg.toy_id, "基础问答", "2", 3), 0.7)
        cfg.rules_snapshot = rules
        results = [{"category": "基础问答", "handler": "basic_qa", "turn_results": [
            {"sheet_row": 2, "case_id": "1", "utterances": [{"user_message": "你好"}], "combined_server_text": "你好", "expected_outcome": "打招呼"},
            {"sheet_row": 3, "case_id": "2", "utterances": [{"user_message": "你好"}], "combined_server_text": "你好", "expected_outcome": "打招呼"}]}]
        answer = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps({"score": 0.9, "pass": True, "is_off_topic": False, "reason": "符合"})))])
        with patch.object(LLMJudge, "_create_client", return_value=(object(), None)), patch.object(LLMJudge, "_chat", return_value=answer):
            post_judge_results(results, cfg)
        self.assertFalse(results[0]["turn_results"][0]["pass"])
        self.assertTrue(results[0]["turn_results"][1]["pass"])
        self.assertEqual(results[0]["turn_results"][0]["judge_result"]["threshold"], 0.99)

    def test_remote_mcp_json_sse_and_two_protocol_versions(self):
        received = []
        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def do_POST(self):
                body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                received.append((body, dict(self.headers)))
                method = body["method"]
                if method == "notifications/initialized":
                    self.send_response(202)
                    self.end_headers()
                    return
                if method == "initialize":
                    result = {"protocolVersion": body["params"]["protocolVersion"], "capabilities": {"tools": {}}}
                else:
                    result = {"tools": [{"name": "remote_ping", "description": "remote", "inputSchema": {"type": "object"}}]}
                data = json.dumps({"jsonrpc": "2.0", "id": body["id"], "result": result}).encode()
                sse = self.path == "/sse"
                if sse:
                    data = b": keepalive\n\ndata: " + data + b"\n\n"
                self.send_response(200)
                self.send_header("Content-Type", "text/event-stream" if sse else "application/json")
                self.send_header("Content-Length", str(len(data)))
                if method == "initialize":
                    self.send_header("Mcp-Session-Id", "test-session")
                self.end_headers()
                self.wfile.write(data)
        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            for version, route in [("2025-03-26", "/json"), ("2025-06-18", "/sse"), ("2025-11-25", "/json"), ("2026-07-28", "/sse")]:
                received.clear()
                source = {"type": "mcp_http", "url": f"http://127.0.0.1:{server.server_port}{route}", "protocol_version": version,
                          "token_env": "TEST_MCP_TOKEN"}
                with patch.dict("os.environ", {"TEST_MCP_TOKEN": "local-test-authorization"}):
                    tools, origin = discover_tools(source, Path("unused"))
                self.assertEqual(tools[0]["name"], "remote_ping")
                self.assertEqual(origin["type"], "mcp_http")
                self.assertEqual(received[-1][1]["Authorization"], "Bearer local-test-authorization")
                self.assertNotIn("local-test-authorization", json.dumps(origin))
                if version == "2026-07-28":
                    self.assertEqual(len(received), 1)
                    self.assertEqual(received[0][0]["params"]["_meta"]["io.modelcontextprotocol/protocolVersion"], version)
                else:
                    self.assertEqual([r[0]["method"] for r in received], ["initialize", "notifications/initialized", "tools/list"])
                    self.assertEqual(received[-1][1]["Mcp-Session-Id"], "test-session")
        finally:
            server.shutdown()
            server.server_close()
            thread.join()

    def test_remote_failure_does_not_fall_back(self):
        with patch("tools_source.requests.Session.post", side_effect=ConnectionError("offline")):
            with self.assertRaises(ConnectionError):
                discover_tools({"type": "mcp_http", "url": "http://127.0.0.1:1/mcp"}, config().mcp_tools_file)

    def test_same_device_file_lock_releases_after_failure(self):
        with device_lock("offline-test-device"):
            with self.assertRaises(ValueError):
                with device_lock("offline-test-device"):
                    pass
        with device_lock("offline-test-device"):
            pass

    def test_event_executes_main_business_once_and_keeps_report_if_notify_fails(self):
        from text_input_test import run_framework
        with tempfile.TemporaryDirectory() as tmp:
            cfg = config()
            cfg.reports_dir = Path(tmp) / "reports"
            cfg.api_key = cfg.deepseek_api_key = cfg.token = "offline-test-key"
            cfg.response_collect_timeout = 0.1
            cfg.send_interval = 0
            cfg.hello_drain_per_wait = 0.001
            cfg.hello_drain_idle_rounds = 1
            cfg.collect_quiet_seconds = 0.001

            def parse(arguments):
                selected = arguments[arguments.index("--case-ids") + 1] if "--case-ids" in arguments else ""
                return replace(cfg, selected_case_ids=selected.split(",") if selected else []), False

            async def connect(*args, **kwargs):
                return FakeSocket()
            response = SimpleNamespace(ok=True, status_code=200, json=lambda: {"websocket": {"url": "ws://test.invalid"}})
            answer = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps({"score": 0.9, "is_off_topic": False, "reason": "符合"})))])
            event = {"event_id": "test-once", "toy_id": "default", "case_ids": ["1"], "notify": True}
            with patch("trigger_run.parse_args", side_effect=parse), patch("collection.requests.post", return_value=response), \
                 patch("collection.websockets.connect", side_effect=connect), patch.object(LLMJudge, "_create_client", return_value=(object(), None)), \
                 patch.object(LLMJudge, "_chat", return_value=answer), patch.dict("os.environ", {"WECOM_WEBHOOK_URL": ""}), \
                 patch("trigger_run.run_framework", wraps=run_framework) as run:
                first = execute_event(event, Path(tmp) / "state")
                second = execute_event(event, Path(tmp) / "state")
            self.assertEqual(first["status"], "completed")
            self.assertEqual(first["notification"], "failed")
            self.assertTrue(second["duplicate"])
            self.assertEqual(run.call_count, 1)
            self.assertTrue(Path(first["report_paths"]["xlsx_export"]).is_file())
            self.assertEqual(xlsx_detailed_stats(Path(first["report_paths"]["xlsx_export"])), (1, 1, 0, 0, 0))
            raw = json.loads(Path(first["report_paths"]["detail_report"]).read_text())
            self.assertEqual(raw["protocol_info"]["trigger_type"], "event")
            self.assertEqual(raw["protocol_info"]["selection"]["case_ids"], ["1"])

    def test_daily_clock_and_duplicate_suppression(self):
        cfg = {"daily_enabled": True, "daily_time": "10:00", "timezone": "Asia/Shanghai", "toy_ids": ["default"], "watch_events": True}
        with tempfile.TemporaryDirectory() as tmp, patch("trigger_run.subprocess.run", return_value=SimpleNamespace(returncode=0)) as run:
            early = datetime(2026, 10, 27, 9, 59, tzinfo=ZoneInfo("Asia/Shanghai"))
            due = early.replace(hour=10, minute=0)
            self.assertIsNone(run_daily_if_due(cfg, tmp, early))
            self.assertEqual(run_daily_if_due(cfg, tmp, due)["status"], "completed")
            self.assertIsNone(run_daily_if_due(cfg, tmp, due))
            self.assertEqual(run.call_count, 1)
            self.assertEqual(run.call_args.kwargs["env"]["TOY_IDS"], "default")

    def test_invalid_event_and_broken_inbox_are_reported(self):
        with self.assertRaises(ValueError):
            validate_event({"event_id": "../escape", "toy_id": "default"})
        with self.assertRaises(ValueError):
            validate_event({"event_id": "valid", "toy_id": "default", "token": "secret"})
        with tempfile.TemporaryDirectory() as tmp:
            inbox = Path(tmp) / "inbox"
            inbox.mkdir()
            (inbox / "bad.json").write_text("broken")
            result = scan_events(tmp)
            self.assertEqual(result[0]["status"], "failed")
            self.assertTrue((Path(tmp) / "invalid/bad.json").is_file())


class CompletionUITests(unittest.TestCase):
    def test_missing_rules_and_invalid_trigger_config_do_not_crash_page(self):
        from streamlit.testing.v1 import AppTest
        from ui_jobs import ROOT
        from config import Config
        with tempfile.TemporaryDirectory() as tmp:
            for field in ["judge_template_file", "thresholds_file"]:
                def factory(**kwargs):
                    return replace(Config(**kwargs), **{field: Path(tmp) / "missing"})
                with patch("config.Config", side_effect=factory):
                    app = AppTest.from_file(str(ROOT / "streamlit_app.py")).run(timeout=60)
                    app.sidebar.radio[0].set_value("评测配置与执行").run()
                self.assertFalse(app.exception)
                self.assertTrue(any("评分规则配置无效" in e.value for e in app.error))
                self.assertTrue(next(b for b in app.button if "开始评测" in b.label).disabled)
            invalid = Path(tmp) / "triggers.json"
            invalid.write_text("broken")
            with patch("ui_management.CONFIG", invalid):
                app = AppTest.from_file(str(ROOT / "streamlit_app.py")).run(timeout=60)
                app.sidebar.radio[0].set_value("评测配置与执行").run()
            self.assertFalse(app.exception)
            self.assertTrue(any("触发配置不可用" in e.value for e in app.error))

    def test_switching_two_reports_updates_metrics(self):
        from streamlit.testing.v1 import AppTest
        from ui_jobs import ROOT
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "reports/default"
            path.mkdir(parents=True)
            for index, passed in enumerate([True, False]):
                data = {"protocol_info": {"toy_id": "default"}, "results": [{"category": "基础问答", "turn_results": [{
                    "sheet_row": 2, "case_id": "1", "pass": passed, "expected_outcome": "打招呼", "combined_server_text": "你好",
                    "judge_result": {"judge_available": True, "score": 0.9 if passed else 0.1, "threshold": 0.7, "reason": "测试"}}]}]}
                file = path / f"default_text_result_{index}.json"
                file.write_text(json.dumps(data))
                os.utime(file, (1000 + index, 1000 + index))
            with patch("ui_jobs.ROOT", Path(tmp)):
                app = AppTest.from_file(str(ROOT / "streamlit_app.py")).run(timeout=60)
                app.sidebar.radio[0].set_value("结果分析").run()
                self.assertEqual(app.metric[1].value, "0")
                next(s for s in app.selectbox if s.label == "选择评测报告").set_value(1).run()
                self.assertEqual(app.metric[1].value, "1")
            self.assertFalse(app.exception)

    def test_ui_management_saves_cases_source_and_case_thresholds(self):
        from streamlit.testing.v1 import AppTest
        from ui_jobs import ROOT
        import streamlit
        from config import Config
        cfg = config()
        original_hash = digest(cfg.cases_xlsx)
        with tempfile.TemporaryDirectory(dir=ROOT, prefix=".test-ui-") as tmp:
            root = Path(tmp)
            toy_root = root / "default"
            toy_root.mkdir()
            path = toy_root / "cases.xlsx"
            shutil.copy2(cfg.cases_xlsx, path)
            profile = toy_root / "toy.json"
            profile.write_text(json.dumps({"token": "", "cases_xlsx": str(path), "mcp_tools_file": str(cfg.mcp_tools_file)}))
            threshold_path = root / "thresholds.json"
            threshold_path.write_text('{"thresholds": {}, "case_thresholds": []}')

            def factory(**kwargs):
                return replace(Config(**kwargs), thresholds_file=threshold_path)

            native_editor = streamlit.data_editor
            def edit_cases(frame, **kwargs):
                if "源行号" in frame.columns:
                    frame = frame.copy()
                    frame.loc[0, "预期结果"] = "通过界面修改的预期"
                    return frame
                return native_editor(frame, **kwargs)

            def edit_thresholds(frame, **kwargs):
                if "启用分类覆盖" in frame.columns:
                    frame = frame.copy()
                    frame.loc[frame["类型"] == "基础问答", "启用分类覆盖"] = True
                    frame.loc[frame["类型"] == "基础问答", "通过阈值"] = 0.8
                    return frame
                if "toy_id" in frame.columns:
                    return pd.DataFrame([{"toy_id": "default", "category": "基础问答", "case_id": "1", "sheet_row": 2, "threshold": 0.95}])
                return native_editor(frame, **kwargs)

            with patch("toy_profiles.TOYS_DIR", root), patch("config.Config", side_effect=factory):
                app = AppTest.from_file(str(ROOT / "streamlit_app.py")).run(timeout=60)
                app.sidebar.radio[0].set_value("用例与工具管理").run()
                with patch("streamlit.data_editor", side_effect=edit_cases):
                    next(b for b in app.button if b.label == "校验并保存用例").click().run()
                self.assertFalse(app.exception)
                self.assertEqual(editable_cases(path)[0]["预期结果"], "通过界面修改的预期")
                next(s for s in app.selectbox if s.label == "来源类型").set_value("mcp_http")
                next(t for t in app.text_input if t.label == "MCP 服务地址").set_value("http://127.0.0.1:1/mcp")
                next(b for b in app.button if b.label == "保存工具来源").click().run()
                self.assertFalse(app.exception)
                self.assertEqual(json.loads(profile.read_text())["tools_source"]["type"], "mcp_http")
                app.sidebar.radio[0].set_value("评测配置与执行").run()
                with patch("streamlit.data_editor", side_effect=edit_thresholds):
                    next(b for b in app.button if b.label == "保存分类与单条阈值").click().run()
                self.assertFalse(app.exception)
                saved = json.loads(threshold_path.read_text())
                self.assertEqual(saved["thresholds"]["基础问答"], 0.8)
                self.assertEqual(saved["case_thresholds"][0]["threshold"], 0.95)
                self.assertTrue(list((toy_root / ".backups").iterdir()))
        self.assertEqual(digest(cfg.cases_xlsx), original_hash)


if __name__ == "__main__":
    unittest.main()
