#!/usr/bin/env python3
"""测试报告：JSON / Markdown / 在原用例表后追加结果列。"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl

from config import _dialogue_worksheet, _ensure_test_result_columns, ensure_dir

# 测试报告生成
class ReportWriter:
    def __init__(self, reports_dir: Path):
        self.reports_dir = reports_dir

    @staticmethod
    def _turn_export_fields(turn_result: Dict[str, Any]) -> Dict[str, Any]:
        """从单条 turn 结果提取写入 xlsx 的字段。"""
        cst = turn_result.get("combined_server_text")
        utterances = turn_result.get("utterances") or []
        # 知识库需要进行两次请求 两次请求分别记录
        if bool(turn_result.get("knowledge_base_dual_round")):
            t1 = str(turn_result.get("kb_text_candidate_1", "") or "").strip()
            t2 = str(turn_result.get("kb_text_candidate_2", "") or "").strip()
            if not t1 and isinstance(turn_result.get("kb_round_1"), dict):
                t1 = str((turn_result["kb_round_1"] or {}).get("combined_server_text", "") or "").strip()
            if not t2 and isinstance(turn_result.get("kb_round_2"), dict):
                t2 = str((turn_result["kb_round_2"] or {}).get("combined_server_text", "") or "").strip()
            actual = f"[第1次回复]\n{t1}\n[第2次回复]\n{t2}".strip()
        # 优先使用 combined_server_text，避免与 utterances 同源字段重复拼接。
        elif isinstance(cst, str) and cst.strip():
            actual = cst.strip()
        else:
            parts: List[str] = []
            for u in utterances:
                if not isinstance(u, dict):
                    continue
                seg = ""
                if isinstance(u.get("server_text_for_judge"), str) and u["server_text_for_judge"].strip():
                    seg = u["server_text_for_judge"].strip()
                elif isinstance(u.get("server_text"), str) and u["server_text"].strip():
                    seg = u["server_text"].strip()
                elif isinstance(u.get("server_actual_when_no_valid_mcp_call"), str) and u["server_actual_when_no_valid_mcp_call"].strip():
                    seg = u["server_actual_when_no_valid_mcp_call"].strip()
                elif isinstance(u.get("assistant_text_from_logs"), str) and u["assistant_text_from_logs"].strip():
                    seg = u["assistant_text_from_logs"].strip()
                elif isinstance(u.get("assistant_visible_text"), str) and u["assistant_visible_text"].strip():
                    seg = u["assistant_visible_text"].strip()
                if seg:
                    parts.append(seg)
            actual = "\n".join(parts) if parts else ""

        mcp_calls: List[Dict[str, Any]] = []
        for u in utterances:
            if not isinstance(u, dict):
                continue
            for c in u.get("mcp_calls") or []:
                if isinstance(c, dict):
                    mcp_calls.append(c)
            mcp_chk = u.get("mcp_check")
            if isinstance(mcp_chk, dict):
                for c in mcp_chk.get("mcp_calls") or []:
                    if isinstance(c, dict):
                        mcp_calls.append(c)

        briefs: List[str] = []
        for c in mcp_calls:
            name = str(c.get("tool_name", "") or "")
            args = c.get("tool_args", {}) or {}
            briefs.append(f"{name}({json.dumps(args, ensure_ascii=False)})")
        mcp_str = "; ".join(briefs) if briefs else ""

        # 普通问答、情感、长期记忆、知识库使用 judge_result；
        # 短期记忆使用 memory_judge。MCP 必调用用例采用规则校验，不经过 LLM 裁判。
        judge_result = turn_result.get("judge_result")
        if not isinstance(judge_result, dict):
            judge_result = turn_result.get("memory_judge")
        if not isinstance(judge_result, dict):
            judge_result = {}

        if judge_result:
            if bool(judge_result.get("judge_pending", False)):
                judge_status = "待评分"
            elif bool(judge_result.get("judge_available", False)):
                judge_status = "正常"
            else:
                judge_status = "评分失败"
        else:
            judge_status = "未使用（MCP规则校验）"

        qwen_score, qwen_reason, deepseek_score, deepseek_reason = self._split_judge_columns(judge_result)
        combined_score = self._combined_score_text(judge_result)

        mcp_reason = ""
        if not judge_result:
            # MCP 规则校验没有 judge_result，改为汇总逐次工具校验原因。
            reason_parts: List[str] = []
            for u in utterances:
                if not isinstance(u, dict):
                    continue
                reason = str(u.get("reason", "") or "").strip()
                if not reason and isinstance(u.get("mcp_check"), dict):
                    reason = str((u.get("mcp_check") or {}).get("reason", "") or "").strip()
                if reason and reason not in reason_parts:
                    reason_parts.append(reason)
            mcp_reason = "; ".join(reason_parts)

        return {
            "pass": bool(turn_result.get("pass", False)),
            "actual": actual,
            "has_mcp": len(mcp_calls) > 0,
            "mcp_params": mcp_str,
            "judge_status": judge_status,
            "qwen_score": qwen_score,
            "qwen_reason": qwen_reason or mcp_reason,
            "deepseek_score": deepseek_score,
            "deepseek_reason": deepseek_reason,
            "combined_score": combined_score,
            "result_label": "评分失败" if judge_status == "评分失败" else ("通过" if bool(turn_result.get("pass", False)) else "不通过"),
        }

    @staticmethod
    def _format_one_judge_score(raw: Dict[str, Any]) -> Any:
        if raw.get("score") is not None:
            return raw.get("score")
        score_parts: List[str] = []
        if raw.get("score_1") is not None:
            score_parts.append(f"第1轮={raw.get('score_1')}")
        if raw.get("score_2") is not None:
            score_parts.append(f"第2轮={raw.get('score_2')}")
        if raw.get("diversity_score") is not None:
            score_parts.append(f"多样性={raw.get('diversity_score')}")
        return "; ".join(score_parts)

    @classmethod
    def _split_judge_columns(cls, judge_result: Dict[str, Any]) -> tuple[Any, str, Any, str]:
        """把合并后的裁判结果拆成 Qwen / DeepSeek 各自的评分和原因。"""
        per_judge = judge_result.get("judges")
        if not isinstance(per_judge, list) or not per_judge:
            if not judge_result:
                return "", "", "", ""
            name = str(judge_result.get("judge_name") or "").strip().lower()
            score = cls._format_one_judge_score(judge_result)
            reason = str(judge_result.get("reason", "") or "").strip()
            if name == "deepseek":
                return "", "", score, reason
            return score, reason, "", ""

        by_name: Dict[str, Dict[str, Any]] = {}
        for raw in per_judge:
            if not isinstance(raw, dict):
                continue
            name = str(raw.get("judge_name") or "").strip().lower()
            if not name:
                continue
            by_name[name] = raw
        qwen = by_name.get("qwen") or {}
        deepseek = by_name.get("deepseek") or {}
        return (
            cls._format_one_judge_score(qwen) if qwen else "",
            str(qwen.get("reason", "") or "").strip() if qwen else "",
            cls._format_one_judge_score(deepseek) if deepseek else "",
            str(deepseek.get("reason", "") or "").strip() if deepseek else "",
        )

    @staticmethod
    def _combined_score_text(judge_result: Dict[str, Any]) -> Any:
        if judge_result.get("score") is not None:
            return judge_result.get("score")
        score_parts: List[str] = []
        if judge_result.get("score_1") is not None:
            score_parts.append(f"第1轮={judge_result.get('score_1')}")
        if judge_result.get("score_2") is not None:
            score_parts.append(f"第2轮={judge_result.get('score_2')}")
        if judge_result.get("diversity_score") is not None:
            score_parts.append(f"多样性={judge_result.get('diversity_score')}")
        return "; ".join(score_parts)

    def _build_sheet_row_export_map(self, results: List[Dict[str, Any]]) -> Dict[int, Dict[str, Any]]:
        """用例结果按 Excel 行号 sheet_row 索引（与 load_dialogue_groups 一致）。"""
        out: Dict[int, Dict[str, Any]] = {}
        for group in results:
            tools_list = group.get("mcp_tools_list")
            if not isinstance(tools_list, dict):
                tools_list = {}
            handshake_fields = {
                "tools_list_received": bool(tools_list.get("received", False)),
                "tools_list_tool_count": int(tools_list.get("returned_tool_count", 0) or 0),
                "tools_list_tool_names": ", ".join(
                    str(name) for name in (tools_list.get("returned_tool_names") or []) if str(name).strip()
                ),
            }

            exported_rows: set[int] = set()
            for tr in group.get("turn_results") or []:
                if not isinstance(tr, dict):
                    continue
                sr = tr.get("sheet_row")
                if sr is None:
                    continue
                try:
                    row_idx = int(sr)
                except (TypeError, ValueError):
                    continue
                row_fields = self._turn_export_fields(tr)
                row_fields.update(handshake_fields)
                out[row_idx] = row_fields
                exported_rows.add(row_idx)

            # 连接或 hello 失败时没有 turn_results，也要把握手信息写回对应的用例行。
            case_spec = group.get("case_spec")
            case_turns = case_spec.get("turns") if isinstance(case_spec, dict) else []
            for turn_spec in case_turns or []:
                if not isinstance(turn_spec, dict):
                    continue
                try:
                    row_idx = int(turn_spec.get("sheet_row"))
                except (TypeError, ValueError):
                    continue
                if row_idx in exported_rows:
                    continue
                out[row_idx] = {
                    "pass": bool(group.get("group_pass", False)),
                    "actual": "",
                    "has_mcp": False,
                    "mcp_params": "",
                    "judge_status": "未执行",
                    "qwen_score": "",
                    "qwen_reason": str(group.get("reason", "") or ""),
                    "deepseek_score": "",
                    "deepseek_reason": "",
                    "combined_score": "",
                    "result_label": "不通过",
                    **handshake_fields,
                }
        return out

    def _export_test_results_to_xlsx(self, cases_xlsx: Path, results: List[Dict[str, Any]], out_path: Path) -> Path:
        """在原用例表基础上追加测试结果、MCP 信息和裁判详情。"""
        row_map = self._build_sheet_row_export_map(results)
        wb = openpyxl.load_workbook(cases_xlsx)
        ws = _dialogue_worksheet(wb)
        removed_report_headers = {
            "mcp调用参数",
            "评分器状态",
            "是否收到tools/list",
            "返回工具数量",
            "返回工具名称",
            "评分",
            "原因",
        }
        # 兼容把旧版导出报告继续作为输入的情况：按列号倒序删除，避免列移动造成漏删。
        for c in range(ws.max_column, 0, -1):
            header = ws.cell(row=1, column=c).value
            if header is not None and str(header).strip() in removed_report_headers:
                ws.delete_cols(c, 1)

        extra_headers = [
            "测试结果是否通过",
            "实际返回内容",
            "是否有mcp调用",
            "Qwen评分",
            "Qwen原因",
            "DeepSeek评分",
            "DeepSeek原因",
            "综合评分",
        ]
        base_col = _ensure_test_result_columns(ws, extra_headers)
        for r in range(2, ws.max_row + 1):
            exp = row_map.get(r)
            if not exp:
                continue
            ws.cell(row=r, column=base_col, value=exp.get("result_label") or ("通过" if exp["pass"] else "不通过"))
            ws.cell(row=r, column=base_col + 1, value=exp["actual"])
            ws.cell(row=r, column=base_col + 2, value="是" if exp["has_mcp"] else "否")
            ws.cell(row=r, column=base_col + 3, value=exp.get("qwen_score", ""))
            ws.cell(row=r, column=base_col + 4, value=exp.get("qwen_reason", ""))
            ws.cell(row=r, column=base_col + 5, value=exp.get("deepseek_score", ""))
            ws.cell(row=r, column=base_col + 6, value=exp.get("deepseek_reason", ""))
            ws.cell(row=r, column=base_col + 7, value=exp.get("combined_score", ""))
        ensure_dir(out_path.parent)
        wb.save(out_path)
        wb.close()
        return out_path

    def write(
        self,
        results: List[Dict[str, Any]],
        protocol_info: Dict[str, Any],
        cases_xlsx: Optional[Path] = None,
        export_xlsx: Optional[Path] = None,
    ) -> Dict[str, str]:
        ensure_dir(self.reports_dir)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        toy_id = str(protocol_info.get("toy_id") or "").strip() if isinstance(protocol_info, dict) else ""
        prefix = f"{toy_id}_" if toy_id else ""
        detail = self.reports_dir / f"{prefix}text_result_{ts}.json"
        summary = self.reports_dir / f"{prefix}text_summary_{ts}.md"

        detail.write_text(
            json.dumps({"protocol_info": protocol_info, "results": results}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        lines = [
            "# 文本输入测试框架摘要",
            "",
            f"- 玩具: {protocol_info.get('toy_id') or '-'}",
            f"- 用例总数: {len(results)}",
            f"- 生成时间: {ts}",
            "",
            "## 说明",
            "- 协议骨架已对齐：OTA -> WebSocket -> hello -> listen/manual/detect",
            "- 裁判模型与评测流程全仓库共用；各玩具只换用例、工具描述和 token",
        ]
        summary.write_text("\n".join(lines) + "\n", encoding="utf-8")
        out: Dict[str, str] = {"detail_report": str(detail), "summary_report": str(summary)}
        if cases_xlsx is not None and cases_xlsx.is_file():
            xlsx_dest = export_xlsx if export_xlsx is not None else (self.reports_dir / f"{prefix}text_result_{ts}_export.xlsx")
            self._export_test_results_to_xlsx(cases_xlsx, results, xlsx_dest)
            out["xlsx_export"] = str(xlsx_dest)
        return out
